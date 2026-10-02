from __future__ import annotations

from pathlib import Path
from shutil import copy2

from click.testing import CliRunner
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]

from bom_monitor_builder.cli import main


def test_design_excel_generic_profile_supports_non_aws_layout(tmp_path: Path) -> None:
    input_path = tmp_path / "sample_service_input.xlsx"
    template_path = tmp_path / "sample_service_template.xlsx"
    profile_path = tmp_path / "sample_service_monitor.yml"
    output_path = tmp_path / "sample_service_output.xlsx"

    copy2(Path("profiles/sample_service_monitor.yml"), profile_path)
    create_generic_input(input_path)
    create_generic_template(template_path)

    runner = CliRunner()
    result = runner.invoke(
        main,
        [
            "design-excel",
            "--input",
            str(input_path),
            "--profile",
            str(profile_path),
            "--output",
            str(output_path),
            "--overwrite",
        ],
    )

    assert result.exit_code == 0, result.output
    workbook = load_workbook(output_path, data_only=False)
    assert workbook.sheetnames == ["Settings", "Reference"]
    assert workbook["Reference"].sheet_state == "hidden"

    settings = workbook["Settings"]
    assert settings["A1"].value == "Converted sample_service_input"
    assert settings["A3"].value == "SG01"
    assert settings["B3"].value == "CorePlatform"
    assert settings["C3"].value == "MON01"
    assert settings["D3"].value == "CPU Service"
    assert settings["F3"].value == "5m"
    assert settings["G3"].value == "有効"
    assert settings["H3"].value == "70以上"
    assert settings["I3"].value == "85以上"
    assert settings["J3"].value == "ops-team"
    assert settings["K3"].value == "none"
    assert settings["L3"].value is None

    assert settings["A4"].value == "SG01"
    assert settings["C4"].value == "MON02"
    assert settings["G4"].value == "無効"
    assert settings["A5"].value == "SG01"
    assert settings["C5"].value == "MON03"

    assert settings["A3"].fill.fill_type == "solid"
    assert settings["A3"].fill.fgColor.rgb == "00D9EAF7"
    assert settings.row_dimensions[3].height == 24
    assert settings.row_dimensions[4].height == 24
    assert settings.column_dimensions["B"].width == 22


def test_generic_profile_does_not_interfere_with_aws_profile(tmp_path: Path) -> None:
    output_path = tmp_path / "aws_output.xlsx"
    runner = CliRunner()

    result = runner.invoke(
        main,
        [
            "design-excel",
            "--input",
            "input/0104_AWS月コスト監視.xlsx",
            "--profile",
            "profiles/aws_cost.yml",
            "--output",
            str(output_path),
            "--overwrite",
        ],
    )

    assert result.exit_code == 0, result.output
    workbook = load_workbook(output_path)
    assert workbook["監視設定"]["C4"].value == "AWSコスト管理"


def create_generic_input(path: Path) -> None:
    workbook = Workbook()
    groups = workbook.active
    groups.title = "Group Catalog"
    groups.append(["Enabled", "Team", "Folder", "Memo"])
    groups.append(["true", "CorePlatform", "SG01", None])

    monitors = workbook.create_sheet("Monitor Catalog")
    monitors.append(["Warn Expr", "Service Label", "Cycle", "Use Flag", "Monitor File", "Kind", "Team", "Crit Expr"])
    monitors.append(["70 以上", "CPU Service", "5 m", "true", "MON01.xml", "Service", "CorePlatform", "85 以上"])
    monitors.append(["60 以上", "Memory Service", "10 m", "false", "MON02.xml", "Service", "CorePlatform", "80 以上"])
    monitors.append(["50 以上", "Disk Service", "15 m", "true", "MON03.xml", "Service", "CorePlatform", "75 以上"])

    details = workbook.create_sheet("Monitor Details")
    details.append(["監視No.", 1])
    details.append(["XMLファイルパス", "Monitor/SG01/MON01.xml"])
    details.append(["OptionalField", "X"])
    workbook.save(path)


def create_generic_template(path: Path) -> None:
    workbook = Workbook()
    settings = workbook.active
    settings.title = "Settings"
    settings["A1"] = "Template Title"
    headers = ["Group ID", "Group Name", "Monitor ID", "Monitor Name", "Type", "Interval", "Enabled", "Warn", "Crit", "Owner", "Notify", "Comment"]
    for column_index, header in enumerate(headers, start=1):
        settings.cell(2, column_index, header)
    values = ["GX", "TemplateGroup", "MONX", "TemplateMonitor", "TypeX", "1m", "有効", "10以上", "20以上", "owner", "notify", "comment"]
    fill = PatternFill(fill_type="solid", fgColor="00D9EAF7")
    for column_index, value in enumerate(values, start=1):
        cell = settings.cell(3, column_index, value)
        cell.fill = fill
        cell.font = Font(bold=(column_index == 1))
    settings.row_dimensions[3].height = 24
    settings.column_dimensions["B"].width = 22
    reference = workbook.create_sheet("Reference")
    reference.sheet_state = "hidden"
    reference["A1"] = "Reference data"
    workbook.save(path)
