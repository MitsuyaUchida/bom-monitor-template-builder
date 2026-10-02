from __future__ import annotations

import json
from pathlib import Path

from click.testing import CliRunner
from openpyxl import load_workbook  # type: ignore[import-untyped]

from bom_monitor_builder.cli import main


def test_activeimage_protector2022_serveredition_build_auto_detects_profile(tmp_path: Path) -> None:
    output_path = tmp_path / "activeimage_protector2022_serveredition_design.xlsx"
    work_dir = tmp_path / "work"
    dump_model_path = tmp_path / "activeimage_protector2022_serveredition_model.json"
    runner = CliRunner()

    result = runner.invoke(
        main,
        [
            "build",
            "--input",
            "import/cab/0905_ActiveImage Protector 2022 ServerEditon.CAB",
            "--output",
            str(output_path),
            "--work-dir",
            str(work_dir),
            "--dump-model",
            str(dump_model_path),
            "--overwrite",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "Detected profile: activeimage_protector2022_serveredition" in result.output
    assert "Template: generated automatically" in result.output
    assert "Groups: 1" in result.output
    assert "Monitors: 4" in result.output
    assert "Validation: PASS" in result.output
    assert output_path.exists()

    model_payload = json.loads(dump_model_path.read_text(encoding="utf-8"))
    assert len(model_payload["groups"]) == 1
    assert len(model_payload["monitors"]) == 4
    assert {monitor["monitor_type"] for monitor in model_payload["monitors"]} == {"Service", "EventlogWSA"}

    workbook = load_workbook(output_path, data_only=False)
    assert workbook.sheetnames == ["環境", "監視設定", "チェックシート"]
    settings = workbook["監視設定"]
    assert settings["B4"].value == "GRP01"
    assert settings["C4"].value == "ActiveImage Protector 2022 ServerEditon"
    assert settings["D4"].value == "MON01"
    assert settings["E4"].value == "ActiveImage Protector Service 監視"
    assert settings["F4"].value == "サービス監視"
    assert settings["K4"].value == "AipService"


def test_activeimage_protector2022_serveredition_design_excel_dry_run_succeeds(tmp_path: Path) -> None:
    parsed_input_path = tmp_path / "0905_ActiveImage_2022_ServerEditon.xlsx"
    dump_model_path = tmp_path / "activeimage_protector2022_serveredition_model.json"
    runner = CliRunner()

    cab_result = runner.invoke(
        main,
        [
            "cab-excel",
            "import/cab/0905_ActiveImage Protector 2022 ServerEditon.CAB",
            "--output",
            str(parsed_input_path),
            "--overwrite",
        ],
    )
    assert cab_result.exit_code == 0, cab_result.output

    dry_run_result = runner.invoke(
        main,
        [
            "design-excel",
            "--input",
            str(parsed_input_path),
            "--profile",
            "profiles/activeimage_protector2022_serveredition.yml",
            "--dump-model",
            str(dump_model_path),
            "--dry-run",
        ],
    )

    assert dry_run_result.exit_code == 0, dry_run_result.output
    payload = json.loads(dry_run_result.output)
    assert payload["profile_id"] == "activeimage_protector2022_serveredition"
    assert payload["template_source"] == "generated_template"
    assert payload["group_count"] == 1
    assert payload["monitor_count"] == 4
