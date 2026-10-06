from pathlib import Path

from click.testing import CliRunner
from openpyxl import load_workbook  # type: ignore[import-untyped]

from bom_monitor_builder.build.detector import analyze_profiles
from bom_monitor_builder.build.exceptions import ProfileDetectionError
from bom_monitor_builder.build.service import build_detection_input
from bom_monitor_builder.cab_excel.cli import parse_input_to_model
from bom_monitor_builder.cli import main


def test_unmatched_cab_uses_generic_and_preserves_actions(tmp_path: Path) -> None:
    source = Path("import/cab/test.CAB")
    parsed, session = parse_input_to_model(source, extracted=False, keep_extracted=False)
    try:
        scores = analyze_profiles(build_detection_input(parsed))
        assert scores
        assert all(candidate.final_score == 0 for candidate in scores)
        assert len(parsed.groups) == 4
        assert len(parsed.items) == 12
        assert len(parsed.actions) == 2
        assert {action.raw_values["Type"] for action in parsed.actions} == {"SendMail"}
    finally:
        session.cleanup()

    output = tmp_path / "generic.xlsx"
    result = CliRunner().invoke(
        main,
        ["build", "--input", str(source), "--output", str(output), "--overwrite", "--verbose"],
    )
    assert result.exit_code == 0, result.output
    assert "Detected profile: generic" in result.output
    assert "Detection mode: generic-fallback" in result.output
    assert "Groups: 4" in result.output
    assert "Monitors: 12" in result.output
    assert "Actions: 2" in result.output
    assert "Validation: PASS" in result.output
    workbook = load_workbook(output, data_only=True)
    assert any(sheet.max_column >= 14 and sheet.max_row >= 15 for sheet in workbook.worksheets)
    assert workbook["Monitor XML"].max_row > 12
    assert workbook["Actions"].max_row == 3


def test_explicit_profile_takes_precedence_over_generic_fallback(tmp_path: Path) -> None:
    output = tmp_path / "explicit.xlsx"
    result = CliRunner().invoke(
        main,
        ["build", "--input", "import/cab/test.CAB", "--profile", "hyperv", "--output", str(output), "--overwrite"],
    )
    assert result.exit_code == 0, result.output
    assert "Detected profile: hyperv" in result.output


def test_ambiguous_automatic_detection_uses_generic_fallback(tmp_path: Path, monkeypatch) -> None:
    def ambiguous_detection(*args, **kwargs):
        raise ProfileDetectionError("Profile detection is ambiguous.")

    monkeypatch.setattr("bom_monitor_builder.build.service.detect_profile", ambiguous_detection)
    output = tmp_path / "ambiguous.xlsx"
    result = CliRunner().invoke(
        main,
        ["build", "--input", "import/cab/test.CAB", "--output", str(output), "--overwrite", "--verbose"],
    )
    assert result.exit_code == 0, result.output
    assert "Detected profile: generic" in result.output
    assert "Detection mode: generic-fallback" in result.output
    assert "automatic-detection:ambiguous" in result.output


def test_cab_without_actions_builds_successfully(tmp_path: Path) -> None:
    source = Path("import/cab/0307_sqlserver2022_windows.cab")
    output = tmp_path / "sqlserver.xlsx"
    result = CliRunner().invoke(
        main,
        ["build", "--input", str(source), "--output", str(output), "--overwrite"],
    )
    assert result.exit_code == 0, result.output
    assert "Detected profile: sqlserver2022_windows" in result.output
    assert "Actions: 0" in result.output
    workbook = load_workbook(output, data_only=True)
    assert "Actions" not in workbook.sheetnames
