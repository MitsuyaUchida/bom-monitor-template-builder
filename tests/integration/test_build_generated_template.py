from __future__ import annotations

from pathlib import Path

from click.testing import CliRunner
from openpyxl import load_workbook  # type: ignore[import-untyped]

from bom_monitor_builder.cli import main


def test_build_nec_esmpro23_uses_generated_template(tmp_path: Path) -> None:
    output_path = tmp_path / "NEC_ESMPRO23.xlsx"
    work_dir = tmp_path / "work"
    runner = CliRunner()

    result = runner.invoke(
        main,
        [
            "build",
            "--input",
            "import/cab/NEC_ESMPRO23.CAB",
            "--output",
            str(output_path),
            "--work-dir",
            str(work_dir),
            "--overwrite",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "Detected profile: nec_esmpro23" in result.output
    assert "Template: generated automatically" in result.output
    assert "Validation: PASS" in result.output
    assert output_path.exists()

    workbook = load_workbook(output_path, data_only=False)
    assert workbook.sheetnames == ["環境", "監視設定", "チェックシート"]
    assert workbook["チェックシート"]["C6"].value == "generated_template"
    assert workbook["チェックシート"]["C7"].value == 17
    assert workbook["チェックシート"]["C8"].value == 1
    settings = workbook["監視設定"]
    monitor_ids = [settings[f"D{row}"].value for row in range(4, settings.max_row + 1) if settings[f"D{row}"].value not in (None, "")]
    assert len(monitor_ids) == 17
    assert settings["B4"].value == "GRP01"
    assert settings["B5"].value in (None, "")
    assert settings.sheet_view.showGridLines is False
    assert settings.freeze_panes == "A4"
    assert settings.auto_filter.ref == "A3:L4"
    assert settings["A1"].fill.fill_type == "solid"
    assert settings["A1"].fill.fgColor.type == "rgb"
    assert settings["A1"].fill.fgColor.rgb == "FFFFFCCC"
    assert settings["A3"].fill.fgColor.rgb == "FFFFFCCC"
    assert settings["B4"].fill.fgColor.rgb == "FFFFFCCC"
    assert settings["I4"].fill.fgColor.rgb == "FFFFFCCC"
    assert settings["J4"].fill.fgColor.rgb == "FFFFFCCC"
    assert "サービス監視" in {settings[f"F{row}"].value for row in range(4, settings.max_row + 1) if settings[f"F{row}"].value not in (None, "")}
    assert "イベントログ監視" in {settings[f"F{row}"].value for row in range(4, settings.max_row + 1) if settings[f"F{row}"].value not in (None, "")}
    assert "5回連続注意" in {settings[f"J{row}"].value for row in range(4, settings.max_row + 1) if settings[f"J{row}"].value not in (None, "")}
