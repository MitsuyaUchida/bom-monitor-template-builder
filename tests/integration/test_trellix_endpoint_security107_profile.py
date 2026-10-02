from __future__ import annotations

import json
from pathlib import Path

from click.testing import CliRunner
from openpyxl import load_workbook  # type: ignore[import-untyped]

from bom_monitor_builder.cli import main


def test_trellix_endpoint_security107_build_auto_detects_profile(tmp_path: Path) -> None:
    output_path = tmp_path / "trellix_endpoint_security107_design.xlsx"
    work_dir = tmp_path / "work"
    dump_model_path = tmp_path / "trellix_endpoint_security107_model.json"
    runner = CliRunner()

    result = runner.invoke(
        main,
        [
            "build",
            "--input",
            "import/cab/0312_Trellix Endpoint Security 107.CAB",
            "--output",
            str(output_path),
            "--work-dir",
            str(work_dir),
            "--dump-model",
            str(dump_model_path),
            "--overwrite",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "Detected profile: trellix_endpoint_security107" in result.output
    assert "Template: generated automatically" in result.output
    assert "Groups: 1" in result.output
    assert "Monitors: 8" in result.output
    assert "Validation: PASS" in result.output
    assert output_path.exists()

    model_payload = json.loads(dump_model_path.read_text(encoding="utf-8"))
    assert len(model_payload["groups"]) == 1
    assert len(model_payload["monitors"]) == 8
    assert {monitor["monitor_type"] for monitor in model_payload["monitors"]} == {"Service", "EventlogWSA"}
    assert model_payload["groups"][0]["group_name"] == "Trellix Endpoint Security 10.7 監視"

    workbook = load_workbook(output_path, data_only=False)
    assert workbook.sheetnames == ["環境", "監視設定", "チェックシート"]
    settings = workbook["監視設定"]
    assert settings["B4"].value == "GRP01"
    assert settings["C4"].value == "Trellix Endpoint Security 10.7 監視"
    assert settings["D4"].value == "MON01"
    assert settings["E4"].value == "Trellix Agent Backwards Compatibility Service 監視"
    assert settings["F4"].value == "サービス監視"
    assert settings["K11"].value == "Application"


def test_trellix_endpoint_security107_design_excel_dry_run_succeeds(tmp_path: Path) -> None:
    parsed_input_path = tmp_path / "0312_Trellix Endpoint Security 107.xlsx"
    dump_model_path = tmp_path / "trellix_endpoint_security107_model.json"
    runner = CliRunner()

    cab_result = runner.invoke(
        main,
        [
            "cab-excel",
            "import/cab/0312_Trellix Endpoint Security 107.CAB",
            "--output",
            str(parsed_input_path),
            "--overwrite",
        ],
    )
    assert cab_result.exit_code == 0, cab_result.output

    dry_run_result = runner.invoke(
        main,
        [
            "design-excel",
            "--input",
            str(parsed_input_path),
            "--profile",
            "profiles/trellix_endpoint_security107.yml",
            "--dump-model",
            str(dump_model_path),
            "--dry-run",
        ],
    )

    assert dry_run_result.exit_code == 0, dry_run_result.output
    payload = json.loads(dry_run_result.output)
    assert payload["profile_id"] == "trellix_endpoint_security107"
    assert payload["template_source"] == "generated_template"
    assert payload["group_count"] == 1
    assert payload["monitor_count"] == 8
