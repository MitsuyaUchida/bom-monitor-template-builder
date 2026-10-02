from __future__ import annotations

import shutil
from pathlib import Path

from click.testing import CliRunner
from openpyxl import load_workbook  # type: ignore[import-untyped]

from bom_monitor_builder.cli import main

PROJECT_ROOT = Path(__file__).resolve().parents[2]


def test_build_standard_cab_rebuilds_monitor_table_without_stale_sqlserver_rows(tmp_path: Path, monkeypatch) -> None:
    prepare_build_workspace(tmp_path, monkeypatch)
    runner = CliRunner()

    result = runner.invoke(
        main,
        [
            "build",
            "--input",
            "import/cab/standard.cab",
            "--profile",
            "profiles/sqlserver2022_windows.yml",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "Detected profile: sqlserver2022_windows" in result.output
    assert "Groups: 3" in result.output
    assert "Monitors: 11" in result.output
    assert "Validation: PASS" in result.output
    assert f"Output: {Path('output/standard.xlsx')}" in result.output

    output_path = tmp_path / "output/standard.xlsx"
    assert output_path.exists()

    workbook = load_workbook(output_path, data_only=False)
    settings = workbook["監視設定"]
    assert settings["A1"].value in (None, "")
    assert settings["A2"].value in (None, "")
    monitor_ids = [settings[f"D{row}"].value for row in range(5, settings.max_row + 1) if settings[f"D{row}"].value not in (None, "")]
    monitor_names = [settings[f"E{row}"].value for row in range(5, settings.max_row + 1) if settings[f"E{row}"].value not in (None, "")]
    group_ids = [settings[f"B{row}"].value for row in range(5, settings.max_row + 1) if settings[f"B{row}"].value not in (None, "")]
    group_names = [settings[f"C{row}"].value for row in range(5, settings.max_row + 1) if settings[f"C{row}"].value not in (None, "")]

    assert len(monitor_ids) == 11
    assert set(monitor_ids) == {"MON01", "MON02", "MON03", "MON04", "MON05", "MON06"}
    assert group_ids == ["GRP01", "GRP02", "GRP03"]
    assert group_names == ["システム監視", "ログ監視", "サービス監視"]
    assert settings["B5"].value == "GRP01"
    assert settings["C5"].value == "システム監視"
    for row in range(6, 11):
        assert settings[f"B{row}"].value in (None, "")
        assert settings[f"C{row}"].value in (None, "")
    assert settings["B11"].value == "GRP02"
    assert settings["C11"].value == "ログ監視"
    assert settings["B12"].value in (None, "")
    assert settings["C12"].value in (None, "")
    assert settings["B13"].value == "GRP03"
    assert settings["C13"].value == "サービス監視"
    for row in range(14, 16):
        assert settings[f"B{row}"].value in (None, "")
        assert settings[f"C{row}"].value in (None, "")
    assert monitor_names == [
        "プロセッサ処理待ち行列長",
        "プロセッサ監視",
        "メモリ監視",
        "仮想メモリ監視",
        "ディスク処理待ち行列長監視",
        "C ドライブディスク容量監視",
        "システムログ監視",
        "アプリケーションログ監視",
        "Server 監視",
        "Remote Procedure Call (RPC) 監視",
        "Windows Management Instrumentation 監視",
    ]
    assert settings.max_row == 15
    assert not any(name in monitor_names for name in ["SQLBackupToUrl イベント監視", "SQLVDI イベント監視", "DatabaseMail イベント監視"])


def prepare_build_workspace(tmp_path: Path, monkeypatch) -> None:
    for directory in ("import", "profiles", "templates"):
        shutil.copytree(PROJECT_ROOT / directory, tmp_path / directory)
    monkeypatch.chdir(tmp_path)
