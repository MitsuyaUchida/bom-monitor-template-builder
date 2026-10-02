from __future__ import annotations

import json
from pathlib import Path

from click.testing import CliRunner
from openpyxl import load_workbook  # type: ignore[import-untyped]

from bom_monitor_builder.cli import main


def test_linux_report_build_auto_detects_profile_and_generates_output(tmp_path: Path) -> None:
    output_path = tmp_path / "linux_report_design.xlsx"
    work_dir = tmp_path / "work"
    dump_model_path = tmp_path / "linux_report_model.json"
    runner = CliRunner()

    result = runner.invoke(
        main,
        [
            "build",
            "--input",
            "import/cab/1002_Linux_report.CAB",
            "--output",
            str(output_path),
            "--work-dir",
            str(work_dir),
            "--dump-model",
            str(dump_model_path),
            "--overwrite",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "Detected profile: linux_report" in result.output
    assert "Template: generated automatically" in result.output
    assert "Groups: 1" in result.output
    assert "Monitors: 18" in result.output
    assert "Validation: PASS" in result.output
    assert output_path.exists()
    assert dump_model_path.exists()

    model_payload = json.loads(dump_model_path.read_text(encoding="utf-8"))
    assert len(model_payload["groups"]) == 1
    assert len(model_payload["monitors"]) == 18
    assert {monitor["monitor_type"] for monitor in model_payload["monitors"]} == {
        "LinuxCpu",
        "LinuxDisk",
        "LinuxDiskStress",
        "LinuxMemory",
        "LinuxNetwork",
        "LinuxTextlog",
    }

    workbook = load_workbook(output_path, data_only=False)
    assert workbook.sheetnames == ["環境", "監視設定", "チェックシート"]
    settings = workbook["監視設定"]
    assert settings["B4"].value == "GRP01"
    assert settings["C4"].value == "Linux レポート向け監視項目"
    assert settings["F4"].value == "LinuxCpu"
    assert settings["D4"].value == "MON01"
    assert settings["E4"].value == "Linux Idle監視"


def test_linux_report_design_excel_dry_run_with_explicit_profile_succeeds(tmp_path: Path) -> None:
    parsed_input_path = tmp_path / "1002_Linux_report.xlsx"
    dump_model_path = tmp_path / "linux_report_model.json"
    runner = CliRunner()

    cab_result = runner.invoke(
        main,
        [
            "cab-excel",
            "import/cab/1002_Linux_report.CAB",
            "--output",
            str(parsed_input_path),
            "--overwrite",
        ],
    )
    assert cab_result.exit_code == 0, cab_result.output

    dry_run_result = runner.invoke(
        main,
        [
            "design-excel",
            "--input",
            str(parsed_input_path),
            "--profile",
            "profiles/linux_report.yml",
            "--dump-model",
            str(dump_model_path),
            "--dry-run",
        ],
    )

    assert dry_run_result.exit_code == 0, dry_run_result.output
    payload = json.loads(dry_run_result.output)
    assert payload["profile_id"] == "linux_report"
    assert payload["template_source"] == "generated_template"
    assert payload["group_count"] == 1
    assert payload["monitor_count"] == 18
