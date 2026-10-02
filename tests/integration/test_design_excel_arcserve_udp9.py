from __future__ import annotations

import json
from pathlib import Path

from click.testing import CliRunner
from openpyxl import load_workbook  # type: ignore[import-untyped]

from bom_monitor_builder.cli import main


def test_arcserve_udp9_real_cab_roundtrip_matches_reference(tmp_path: Path) -> None:
    cab_path = Path("import/cab/Arcserve_UDP_9_10.CAB")
    profile_path = Path("profiles/arcserve_udp9.yml")
    template_path = Path("templates/ArcserveUDP9設計書.xlsx")
    reference_path = Path("reference/ArcserveUDP9設計書.xlsx")
    parsed_input_path = tmp_path / "0234_Arcserve UDP 9.0 監視.xlsx"
    output_path = tmp_path / "ArcserveUDP9設計書.xlsx"
    dump_model_path = tmp_path / "arcserve_udp9_model.json"
    runner = CliRunner()

    cab_result = runner.invoke(
        main,
        [
            "cab-excel",
            str(cab_path),
            "--output",
            str(parsed_input_path),
            "--overwrite",
        ],
    )
    assert cab_result.exit_code == 0, cab_result.output
    assert parsed_input_path.exists()

    dry_run_result = runner.invoke(
        main,
        [
            "design-excel",
            "--input",
            str(parsed_input_path),
            "--profile",
            str(profile_path),
            "--dump-model",
            str(dump_model_path),
            "--dry-run",
        ],
    )
    assert dry_run_result.exit_code == 0, dry_run_result.output
    dry_run_payload = json.loads(dry_run_result.output)
    assert dry_run_payload["monitor_count"] == 11
    assert dry_run_payload["group_count"] == 1
    assert dry_run_payload["input_unchanged"] is True
    assert dry_run_payload["template_unchanged"] is True

    model_payload = json.loads(dump_model_path.read_text(encoding="utf-8"))
    assert sorted(model_payload.keys()) == ["extensions", "groups", "metadata", "monitors"]
    assert len(model_payload["groups"]) == 1
    assert len(model_payload["monitors"]) == 11
    assert {monitor["monitor_type"] for monitor in model_payload["monitors"]} == {"EventlogWSA", "Service"}
    assert "TEST_PASSWORD_DO_NOT_USE" not in dump_model_path.read_text(encoding="utf-8")

    render_result = runner.invoke(
        main,
        [
            "design-excel",
            "--input",
            str(parsed_input_path),
            "--profile",
            str(profile_path),
            "--template",
            str(template_path),
            "--output",
            str(output_path),
            "--overwrite",
        ],
    )
    assert render_result.exit_code == 0, render_result.output
    render_payload = json.loads(render_result.output)
    assert render_payload["monitor_count"] == 11
    assert render_payload["group_count"] == 1
    assert render_payload["input_unchanged"] is True
    assert render_payload["template_unchanged"] is True
    assert render_payload["validation_checks"] == [
        "sheet order",
        "sheet visibility",
        "template formatting",
        "input workbook readable",
        "mapped values",
        "monitor count",
        "group count",
        "no stale monitor rows",
        "no stale template data",
        "merge ranges",
        "alignment rules",
        "border rules",
        "row height rules",
        "secret masking",
    ]

    validate_result = runner.invoke(
        main,
        [
            "design-excel",
            "--input",
            str(parsed_input_path),
            "--profile",
            str(profile_path),
            "--template",
            str(template_path),
            "--output",
            str(output_path),
            "--validate-only",
        ],
    )
    assert validate_result.exit_code == 0, validate_result.output

    reference_wb = load_workbook(reference_path, data_only=False)
    output_wb = load_workbook(output_path, data_only=False)

    assert output_wb.sheetnames == reference_wb.sheetnames
    assert [ws.sheet_state for ws in output_wb.worksheets] == [ws.sheet_state for ws in reference_wb.worksheets]

    reference_settings = reference_wb["監視設定"]
    output_settings = output_wb["監視設定"]
    assert [str(item) for item in output_settings.merged_cells.ranges] == [str(item) for item in reference_settings.merged_cells.ranges]
    assert output_settings.freeze_panes == reference_settings.freeze_panes
    assert [output_settings.row_dimensions[row].height for row in range(1, 15)] == [
        reference_settings.row_dimensions[row].height for row in range(1, 15)
    ]
    assert [output_settings.column_dimensions[column].width for column in "ABCDEFGHIJKLMNOPQRS"] == [
        reference_settings.column_dimensions[column].width for column in "ABCDEFGHIJKLMNOPQRS"
    ]

    for sheet_name in reference_wb.sheetnames:
        reference_ws = reference_wb[sheet_name]
        output_ws = output_wb[sheet_name]
        assert reference_ws.max_row == output_ws.max_row
        assert reference_ws.max_column == output_ws.max_column
        for row in range(1, reference_ws.max_row + 1):
            for column in range(1, reference_ws.max_column + 1):
                assert output_ws.cell(row, column).value == reference_ws.cell(row, column).value

    assert output_wb["環境"]["B6"].value == 'TEST_PASSWORD_DO_NOT_USE'
