from __future__ import annotations

import json
from pathlib import Path

from click.testing import CliRunner
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]

from bom_monitor_builder.cli import main


def test_design_excel_generates_aws_cost_design_workbook(tmp_path: Path) -> None:
    output_path = tmp_path / "AWSコスト監視設計書.xlsx"
    dump_model_path = tmp_path / "model.json"
    runner = CliRunner()

    result = runner.invoke(
        main,
        [
            "design-excel",
            "--input",
            "input/0104_AWS月コスト監視.xlsx",
            "--profile",
            "profiles/aws_cost.yml",
            "--output",
            str(output_path),
            "--overwrite",
            "--dump-model",
            str(dump_model_path),
        ],
    )

    assert result.exit_code == 0, result.output
    assert output_path.exists()
    assert dump_model_path.exists()
    dumped_model = json.loads(dump_model_path.read_text(encoding="utf-8"))
    assert sorted(dumped_model.keys()) == ["extensions", "groups", "metadata", "monitors"]
    assert isinstance(dumped_model["groups"], list)
    assert isinstance(dumped_model["monitors"], list)
    assert dumped_model["extensions"]["aws_cost"]["details"]["MON01"]["ObjectName"] == "C:\\aws-cost\\print_monthly_total_cost.bat"

    workbook = load_workbook(output_path, data_only=False)
    assert workbook.sheetnames == [
        "動作原理",
        "環境",
        "監視設定",
        "チェックシート",
        "AutoTemplate.ymlチェック",
        "サービス追加方法",
    ]
    assert workbook["チェックシート"].sheet_state == "hidden"
    settings = workbook["監視設定"]
    assert settings["A1"].value == "0999_Windows その他：0104_AWS月コスト監視"
    assert settings["B4"].value == "GRP01"
    assert settings["C4"].value == "AWSコスト管理"
    assert settings["D4"].value == "MON01"
    assert settings["E4"].value == "AWS月コスト監視"
    assert settings["F4"].value == "1日"
    assert settings["G4"].value == "有効"
    assert settings["I4"].value == "500以上"
    assert settings["J4"].value == "900以上"
    assert settings["M4"].value == "act01"
    assert settings["N4"].value == "コスト超過通知メール送信"
    assert settings["P4"].value == "注意、危険、失敗"
    assert "secret_access_key" in workbook["環境"]["A6"].value
    assert "****" in workbook["環境"]["A6"].value


def test_design_excel_supports_new_profile_without_core_changes(tmp_path: Path) -> None:
    input_path = tmp_path / "simple_input.xlsx"
    template_path = tmp_path / "simple_template.xlsx"
    profile_path = tmp_path / "simple.yml"
    output_path = tmp_path / "simple_output.xlsx"
    create_simple_input(input_path)
    create_simple_template(template_path)
    profile_path.write_text(
        """
profile:
  id: simple_monitor
  name: Simple Monitor
  version: 1
template:
  path: simple_template.xlsx
source:
  header_search_rows: 5
  sheets:
    groups:
      candidates: ["Groups"]
    monitors:
      candidates: ["Monitors"]
    details:
      candidates: ["Details"]
      required: false
  fields:
    group_id:
      sections: ["groups"]
      aliases: ["Folder"]
    group_name:
      sections: ["groups", "monitors"]
      aliases: ["Group Name"]
    enabled:
      sections: ["groups", "monitors"]
      aliases: ["Enabled"]
    comment:
      sections: ["groups", "monitors"]
      aliases: ["Comment"]
    monitor_id:
      sections: ["monitors"]
      aliases: ["Monitor File"]
    monitor_name:
      sections: ["monitors"]
      aliases: ["Monitor Name"]
    monitor_type:
      sections: ["monitors"]
      aliases: ["Monitor Type"]
    interval:
      sections: ["monitors"]
      aliases: ["Interval"]
    warning_condition:
      sections: ["monitors"]
      aliases: ["Warn"]
    critical_condition:
      sections: ["monitors"]
      aliases: ["Crit"]
transform:
  replacements: {}
  defaults:
    owner: ops-team
  derived_fields: {}
  normalizers:
    monitor_id: ["strip_extension"]
    enabled: ["normalize_boolean"]
    warning_condition: ["normalize_threshold"]
    critical_condition: ["normalize_threshold"]
  plugins: []
output:
  sheets: {}
  tables:
    settings:
      sheet: "Settings"
      start_row: 2
      template_row: 2
      clear_existing_rows: false
      columns:
        A: group_id
        B: group_name
        C: monitor_id
        D: monitor_name
        E: interval
        F: enabled
        G: warning_condition
        H: critical_condition
        I: owner
  cells:
    title:
      sheet: "Settings"
      cell: "A1"
      value: "Converted {{ source_stem }}"
  preserve: {}
validation:
  required_fields: ["group_id", "monitor_id"]
  unique_fields: ["monitor_id"]
  expected_sheets: ["Settings", "Reference"]
security:
  secret_patterns: ["password", "token"]
  output_policy: mask
""".strip(),
        encoding="utf-8",
    )

    runner = CliRunner()
    result = runner.invoke(
        main,
        [
            "design-excel",
            "--input",
            str(input_path),
            "--profile",
            str(profile_path),
            "--output",
            str(output_path),
            "--overwrite",
        ],
    )

    assert result.exit_code == 0, result.output
    workbook = load_workbook(output_path)
    settings = workbook["Settings"]
    assert settings["A1"].value == "Converted simple_input"
    assert settings["A2"].value == "G01"
    assert settings["B2"].value == "Platform"
    assert settings["C2"].value == "MON01"
    assert settings["F2"].value == "有効"
    assert settings["I2"].value == "ops-team"
    assert settings["A3"].value == "G01"
    assert settings["C3"].value == "MON02"
    assert settings["F3"].value == "無効"


def create_simple_input(path: Path) -> None:
    workbook = Workbook()
    groups = workbook.active
    groups.title = "Groups"
    groups.append(["Folder", "Group Name", "Enabled", "Comment"])
    groups.append(["G01", "Platform", "true", "group comment"])

    monitors = workbook.create_sheet("Monitors")
    monitors.append(["Group Name", "Monitor File", "Monitor Name", "Monitor Type", "Enabled", "Interval", "Warn", "Crit", "Comment"])
    monitors.append(["Platform", "MON01.xml", "CPU", "Perf", "true", "5m", "80 以上", "90 以上", "A"])
    monitors.append(["Platform", "MON02.xml", "Memory", "Perf", "false", "10m", "70 以上", "85 以上", "B"])
    workbook.save(path)


def create_simple_template(path: Path) -> None:
    workbook = Workbook()
    settings = workbook.active
    settings.title = "Settings"
    settings["A1"] = "Template"
    settings["A2"] = "GXX"
    settings["B2"] = "Template Group"
    settings["C2"] = "MONXX"
    settings["D2"] = "Template Monitor"
    settings["E2"] = "1m"
    settings["F2"] = "有効"
    settings["G2"] = "10以上"
    settings["H2"] = "20以上"
    settings["I2"] = "owner"
    settings.row_dimensions[2].height = 22
    settings.column_dimensions["B"].width = 25
    reference = workbook.create_sheet("Reference")
    reference.sheet_state = "hidden"
    reference["A1"] = "token=abcd"
    workbook.save(path)
