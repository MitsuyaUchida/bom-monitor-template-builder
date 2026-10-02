from __future__ import annotations

import json
from pathlib import Path

from click.testing import CliRunner
from openpyxl import load_workbook  # type: ignore[import-untyped]

from bom_monitor_builder.cli import main


def test_msfc_2022_build_auto_detects_profile_and_generates_output(tmp_path: Path) -> None:
    output_path = tmp_path / "msfc_2022_design.xlsx"
    work_dir = tmp_path / "work"
    dump_model_path = tmp_path / "msfc_2022_model.json"
    runner = CliRunner()

    result = runner.invoke(
        main,
        [
            "build",
            "--input",
            "import/cab/msfc-2022.CAB",
            "--output",
            str(output_path),
            "--work-dir",
            str(work_dir),
            "--dump-model",
            str(dump_model_path),
            "--overwrite",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "Detected profile: msfc_2022" in result.output
    assert "Template: generated automatically" in result.output
    assert "Groups: 1" in result.output
    assert "Monitors: 7" in result.output
    assert "Validation: PASS" in result.output
    assert output_path.exists()
    assert dump_model_path.exists()

    model_payload = json.loads(dump_model_path.read_text(encoding="utf-8"))
    assert len(model_payload["groups"]) == 1
    assert len(model_payload["monitors"]) == 7
    assert {monitor["monitor_type"] for monitor in model_payload["monitors"]} == {"EventlogWSA", "Perf", "Service"}
    assert model_payload["groups"][0]["group_name"] == "MSFC_Windows Server 2022  監視"
    assert model_payload["extensions"]["msfc_2022"]["details"]["GRP01/MON01"]["監視名"] == "Cluster Service 監視"
    assert "user ****" in dump_model_path.read_text(encoding="utf-8")

    workbook = load_workbook(output_path, data_only=False)
    assert workbook.sheetnames == ["環境", "監視設定", "チェックシート"]
    assert workbook["環境"]["B8"].value == 1
    assert workbook["環境"]["B9"].value == 7
    assert workbook["チェックシート"]["C6"].value == "generated_template"
    settings = workbook["監視設定"]
    assert settings["B4"].value == "GRP01"
    assert settings["C4"].value == "MSFC_Windows Server 2022  監視"
    assert settings["D4"].value == "MON01"
    assert settings["E4"].value == "Cluster Service 監視"
    assert settings["F4"].value == "サービス監視"
    assert settings["K4"].value == "ClusSvc"
