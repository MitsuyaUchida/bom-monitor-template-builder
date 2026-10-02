from __future__ import annotations

import shutil
import subprocess
from datetime import UTC, datetime
from pathlib import Path

from click.testing import CliRunner
from openpyxl import load_workbook  # type: ignore[import-untyped]

from bom_monitor_builder.build.codex_service import (
    BuildFailure,
    BuildFailureKind,
    CodexAutoResult,
    CodexPromptBundle,
)
from bom_monitor_builder.cli import main

PROJECT_ROOT = Path(__file__).resolve().parents[2]


def test_build_sqlserver2022_windows_matches_reference(tmp_path: Path) -> None:
    output_path = tmp_path / "sqlserver2022_windows_design.xlsx"
    work_dir = tmp_path / "work"
    dump_model_path = tmp_path / "sqlserver2022_windows_model.json"
    runner = CliRunner()

    result = runner.invoke(
        main,
        [
            "build",
            "--input",
            "import/cab/0307_sqlserver2022_windows.cab",
            "--output",
            str(output_path),
            "--work-dir",
            str(work_dir),
            "--keep-intermediate",
            "--dump-model",
            str(dump_model_path),
            "--overwrite",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "Detected profile: sqlserver2022_windows" in result.output
    assert "Groups: 3" in result.output
    assert "Monitors: 39" in result.output
    assert "Validation: PASS" in result.output
    assert (work_dir / "0307_sqlserver2022_windows.xlsx").exists()
    assert output_path.exists()
    assert dump_model_path.exists()
    assert diff_count(Path("reference/sqlserver2022_windows_design.xlsx"), output_path) == 0


def test_build_sqlserver2022_windows_creates_new_file_without_overwriting(tmp_path: Path, monkeypatch) -> None:
    prepare_build_workspace(tmp_path, monkeypatch)
    runner = CliRunner()

    first_result = runner.invoke(
        main,
        [
                "build",
                "--input",
                "import/cab/0307_sqlserver2022_windows.cab",
            ],
        )
    assert first_result.exit_code == 0, first_result.output
    assert f"Output: {Path('output/0307_sqlserver2022_windows.xlsx')}" in first_result.output

    first_output = tmp_path / "output/0307_sqlserver2022_windows.xlsx"
    assert first_output.exists()
    first_snapshot = first_output.read_bytes()

    second_result = runner.invoke(
        main,
        [
                "build",
                "--input",
                "import/cab/0307_sqlserver2022_windows.cab",
            ],
        )
    assert second_result.exit_code == 0, second_result.output
    assert f"Output: {Path('output/0307_sqlserver2022_windows_001.xlsx')}" in second_result.output
    assert "Validation: PASS" in second_result.output

    second_output = tmp_path / "output/0307_sqlserver2022_windows_001.xlsx"
    assert second_output.exists()
    assert first_output.read_bytes() == first_snapshot
    assert diff_count(PROJECT_ROOT / "reference/sqlserver2022_windows_design.xlsx", first_output) == 0
    assert diff_count(PROJECT_ROOT / "reference/sqlserver2022_windows_design.xlsx", second_output) == 0


def test_build_arcserve_udp9_matches_reference(tmp_path: Path) -> None:
    output_path = tmp_path / "ArcserveUDP9設計書.xlsx"
    work_dir = tmp_path / "work"
    runner = CliRunner()

    result = runner.invoke(
        main,
        [
            "build",
            "--input",
            "import/cab/Arcserve_UDP_9_10.CAB",
            "--output",
            str(output_path),
            "--work-dir",
            str(work_dir),
            "--keep-intermediate",
            "--overwrite",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "Detected profile: arcserve_udp9" in result.output
    assert "Groups: 1" in result.output
    assert "Monitors: 11" in result.output
    assert "Validation: PASS" in result.output


def test_build_arcserve_udp9_creates_new_file_without_overwriting(tmp_path: Path, monkeypatch) -> None:
    prepare_build_workspace(tmp_path, monkeypatch)
    runner = CliRunner()

    first_result = runner.invoke(
        main,
        [
                "build",
                "--input",
                "import/cab/Arcserve_UDP_9_10.CAB",
            ],
        )
    assert first_result.exit_code == 0, first_result.output
    assert f"Output: {Path('output/Arcserve_UDP_9_10.xlsx')}" in first_result.output

    second_result = runner.invoke(
        main,
        [
                "build",
                "--input",
                "import/cab/Arcserve_UDP_9_10.CAB",
            ],
        )
    assert second_result.exit_code == 0, second_result.output
    assert f"Output: {Path('output/Arcserve_UDP_9_10_001.xlsx')}" in second_result.output

    first_output = tmp_path / "output/Arcserve_UDP_9_10.xlsx"
    second_output = tmp_path / "output/Arcserve_UDP_9_10_001.xlsx"
    assert first_output.exists()
    assert second_output.exists()


def test_build_aws_cost_succeeds_and_masks_secrets(tmp_path: Path) -> None:
    output_path = tmp_path / "aws_cost_design.xlsx"
    runner = CliRunner()

    result = runner.invoke(
        main,
        [
            "build",
            "--input",
            "import/cab/0104_AWS月コスト監視.cab",
            "--output",
            str(output_path),
            "--overwrite",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "Detected profile: aws_cost" in result.output
    assert "Groups: 1" in result.output
    assert "Monitors: 1" in result.output
    workbook = load_workbook(output_path, data_only=False)
    assert "secret_access_key" in str(workbook["環境"]["A6"].value)
    assert "****" in str(workbook["環境"]["A6"].value)


def test_build_sqlserver2022_windows_dry_run_does_not_write_output(tmp_path: Path) -> None:
    output_path = tmp_path / "sqlserver2022_windows_design.xlsx"
    work_dir = tmp_path / "work"
    runner = CliRunner()

    result = runner.invoke(
        main,
        [
            "build",
            "--input",
            "import/cab/0307_sqlserver2022_windows.cab",
            "--output",
            str(output_path),
            "--work-dir",
            str(work_dir),
            "--dry-run",
            "--overwrite",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "Detected profile: sqlserver2022_windows" in result.output
    assert "Planned output: " in result.output
    assert "Groups: 3" in result.output
    assert "Monitors: 39" in result.output
    assert "Validation: SKIPPED" in result.output
    assert not output_path.exists()
    assert (work_dir / "0307_sqlserver2022_windows.xlsx").exists()


def test_build_sqlserver2022_windows_supports_explicit_profile_and_template(tmp_path: Path) -> None:
    output_path = tmp_path / "test.xlsx"
    work_dir = tmp_path / "work"
    runner = CliRunner()

    result = runner.invoke(
        main,
        [
            "build",
            "--input",
            "import/cab/0307_sqlserver2022_windows.cab",
            "--profile",
            "profiles/sqlserver2022_windows.yml",
            "--template",
            "templates/sqlserver2022_windows_design.xlsx",
            "--output",
            str(output_path),
            "--work-dir",
            str(work_dir),
            "--overwrite",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "Detected profile: sqlserver2022_windows" in result.output
    assert output_path.exists()
    assert diff_count(Path("reference/sqlserver2022_windows_design.xlsx"), output_path) == 0


def test_build_validate_only_uses_latest_existing_output_when_output_is_omitted(tmp_path: Path, monkeypatch) -> None:
    prepare_build_workspace(tmp_path, monkeypatch)
    runner = CliRunner()

    first_result = runner.invoke(
        main,
        [
                "build",
                "--input",
                "import/cab/0307_sqlserver2022_windows.cab",
            ],
        )
    assert first_result.exit_code == 0, first_result.output

    second_result = runner.invoke(
        main,
        [
                "build",
                "--input",
                "import/cab/0307_sqlserver2022_windows.cab",
            ],
        )
    assert second_result.exit_code == 0, second_result.output

    validate_result = runner.invoke(
        main,
        [
                "build",
                "--input",
                "import/cab/0307_sqlserver2022_windows.cab",
                "--validate-only",
            ],
        )
    assert validate_result.exit_code == 0, validate_result.output
    assert f"Output: {Path('output/0307_sqlserver2022_windows_001.xlsx')}" in validate_result.output
    assert "Validation: PASS" in validate_result.output


def test_build_with_codex_prompt_writes_prompt_and_logs_for_profile_not_found(tmp_path: Path, monkeypatch) -> None:
    runner = CliRunner()
    log_dir = tmp_path / "codex_logs"
    bundle_dir = log_dir / datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ") / "test_case"
    bundle_dir.mkdir(parents=True)
    (bundle_dir / "codex_prompt.txt").write_text(
        "\n".join(
            [
                "No matching profile was found.",
                "既存profiles一覧:",
                "monitor_typeだけの一致を製品固有の根拠にしないこと。",
            ]
        )
        + "\n",
        encoding="utf-8",
    )

    def fake_run_build_with_codex(*args, **kwargs):
        return (
            False,
            (bundle_dir / "codex_prompt.txt").read_text(encoding="utf-8"),
            CodexPromptBundle(
                prompt=(bundle_dir / "codex_prompt.txt").read_text(encoding="utf-8"),
                log_dir=bundle_dir,
                failure=BuildFailure(
                    kind=BuildFailureKind.PROFILE_NOT_FOUND,
                    message="No matching profile was found.",
                    exit_code=1,
                    stdout="",
                    stderr="No matching profile was found.",
                    intermediate_excel=tmp_path / "example.xlsx",
                ),
                summary=None,
                detection_input=None,
                profile_analyses=[],
            ),
        )

    monkeypatch.setattr("bom_monitor_builder.build.cli.run_build_with_codex", fake_run_build_with_codex)

    result = runner.invoke(
        main,
        [
            "build-with-codex",
            "--input",
            "import/cab/0102_Windows_report.cab",
            "--codex-prompt",
            "--codex-log-dir",
            str(log_dir),
        ],
    )

    assert result.exit_code != 0
    assert "Build failed. Codex investigation prompt was generated." in result.output
    assert "Codex log: " in result.output

    prompt_files = sorted(log_dir.rglob("codex_prompt.txt"))
    assert len(prompt_files) == 1
    prompt = prompt_files[0].read_text(encoding="utf-8")
    assert "No matching profile was found." in prompt
    assert "既存profiles一覧:" in prompt
    assert "monitor_typeだけの一致を製品固有の根拠にしないこと。" in prompt


def test_build_with_codex_auto_runs_codex_and_reruns_build(tmp_path: Path, monkeypatch) -> None:
    runner = CliRunner()
    log_dir = tmp_path / "codex_logs"
    bundle_dir = log_dir / datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ") / "test_case"
    bundle_dir.mkdir(parents=True)

    def fake_run_build_with_codex(*args, **kwargs):
        return (
            0,
            "",
            CodexAutoResult(
                prompt_bundle=CodexPromptBundle(
                    prompt="prompt",
                    log_dir=bundle_dir,
                    failure=BuildFailure(
                        kind=BuildFailureKind.PROFILE_NOT_FOUND,
                        message="No matching profile was found.",
                        exit_code=1,
                        stdout="",
                        stderr="No matching profile was found.",
                        intermediate_excel=tmp_path / "example.xlsx",
                    ),
                    summary=None,
                    detection_input=None,
                    profile_analyses=[],
                ),
                codex_exit_code=0,
                codex_stdout="codex ok\n",
                codex_stderr="",
                rerun_exit_code=0,
                rerun_stdout=(
                    "Build completed successfully.\n\n"
                    "Input CAB: import/cab/0102_Windows_report.cab\n"
                    "Detected profile: linux_report\n"
                    "Template: generated automatically\n"
                    f"Intermediate Excel: {tmp_path / 'example.xlsx'}\n"
                    "Output: output/example.xlsx\n"
                    "Groups: 1\n"
                    "Monitors: 18\n"
                    "Validation: PASS\n"
                ),
                rerun_stderr="",
            ),
        )

    monkeypatch.setattr("bom_monitor_builder.build.cli.run_build_with_codex", fake_run_build_with_codex)

    result = runner.invoke(
        main,
        [
            "build-with-codex",
            "--input",
            "import/cab/0102_Windows_report.cab",
            "--codex-auto",
            "--codex-log-dir",
            str(log_dir),
        ],
    )

    assert result.exit_code == 0, result.output
    assert "Build failed and Codex auto-investigation was executed." in result.output
    assert "Codex exit code: 0" in result.output
    assert "Rebuild exit code: 0" in result.output


def prepare_build_workspace(tmp_path: Path, monkeypatch) -> None:
    for directory in ("import", "profiles", "templates"):
        shutil.copytree(PROJECT_ROOT / directory, tmp_path / directory)
    monkeypatch.chdir(tmp_path)


def diff_count(reference_path: Path, output_path: Path) -> int:
    reference_wb = load_workbook(reference_path, data_only=False)
    output_wb = load_workbook(output_path, data_only=False)
    diffs = 0
    if reference_wb.sheetnames != output_wb.sheetnames:
        diffs += 1
    for sheet_name in reference_wb.sheetnames:
        reference_ws = reference_wb[sheet_name]
        output_ws = output_wb[sheet_name]
        for row in range(1, max(reference_ws.max_row, output_ws.max_row) + 1):
            for column in range(1, max(reference_ws.max_column, output_ws.max_column) + 1):
                if reference_ws.cell(row, column).value != output_ws.cell(row, column).value:
                    diffs += 1
    return diffs
