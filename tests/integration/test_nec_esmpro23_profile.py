from __future__ import annotations

import json
from pathlib import Path

from click.testing import CliRunner
from openpyxl import load_workbook  # type: ignore[import-untyped]

from bom_monitor_builder.cli import main


def test_nec_esmpro23_cab_analysis_and_dry_run_succeeds(tmp_path: Path) -> None:
    cab_path = Path("import/cab/NEC_ESMPRO23.CAB")
    profile_path = Path("profiles/nec_esmpro23.yml")
    parsed_input_path = tmp_path / "nec_esmpro23.xlsx"
    dump_model_path = tmp_path / "nec_esmpro23_model.json"
    runner = CliRunner()

    cab_result = runner.invoke(
        main,
        [
            "cab-excel",
            str(cab_path),
            "--output",
            str(parsed_input_path),
            "--overwrite",
        ],
    )

    assert cab_result.exit_code == 0, cab_result.output
    assert parsed_input_path.exists()

    workbook = load_workbook(parsed_input_path, data_only=True)
    assert workbook.sheetnames == ["表紙", "監視グループ一覧", "監視項目一覧", "監視項目詳細", "XML全項目", "解析情報"]

    dry_run_result = runner.invoke(
        main,
        [
            "design-excel",
            "--input",
            str(parsed_input_path),
            "--profile",
            str(profile_path),
            "--dump-model",
            str(dump_model_path),
            "--dry-run",
        ],
    )

    assert dry_run_result.exit_code == 0, dry_run_result.output
    payload = json.loads(dry_run_result.output)
    assert payload["profile_id"] == "nec_esmpro23"
    assert payload["template_path"] is None
    assert payload["template_source"] == "generated_template"
    assert payload["group_count"] == 1
    assert payload["monitor_count"] == 17
    assert payload["input_unchanged"] is True
    assert payload["template_unchanged"] is True

    model_payload = json.loads(dump_model_path.read_text(encoding="utf-8"))
    assert sorted(model_payload.keys()) == ["extensions", "groups", "metadata", "monitors"]
    assert len(model_payload["groups"]) == 1
    assert len(model_payload["monitors"]) == 17
    assert {monitor["monitor_type"] for monitor in model_payload["monitors"]} == {"EventlogWSA", "Service"}
    assert model_payload["groups"][0]["group_name"] == "NEC ESMPRO/ServerAgent Service Ver2.3 監視"
    assert model_payload["extensions"]["nec_esmpro23"]["details"]["GRP01/MON01"]["監視名"] == "Alert Manager Main Service 監視"
    assert "password" not in dump_model_path.read_text(encoding="utf-8").casefold()


def test_nec_esmpro23_build_auto_detects_profile_without_template(tmp_path: Path) -> None:
    dump_model_path = tmp_path / "nec_esmpro23_model.json"
    work_dir = tmp_path / "work"
    runner = CliRunner()

    result = runner.invoke(
        main,
        [
            "build",
            "--input",
            "import/cab/NEC_ESMPRO23.CAB",
            "--work-dir",
            str(work_dir),
            "--dry-run",
            "--dump-model",
            str(dump_model_path),
            "--overwrite",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "Detected profile: nec_esmpro23" in result.output
    assert "Template: generated automatically" in result.output
    assert "Groups: 1" in result.output
    assert "Monitors: 17" in result.output
    assert "Validation: SKIPPED" in result.output
    assert (work_dir / "NEC_ESMPRO23.xlsx").exists()
    assert dump_model_path.exists()


def test_nec_esmpro23_build_generates_output_without_registered_template(tmp_path: Path) -> None:
    work_dir = tmp_path / "work"
    output_path = tmp_path / "NEC_ESMPRO23.xlsx"
    runner = CliRunner()

    result = runner.invoke(
        main,
        [
            "build",
            "--input",
            "import/cab/NEC_ESMPRO23.CAB",
            "--work-dir",
            str(work_dir),
            "--output",
            str(output_path),
            "--overwrite",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "Detected profile: nec_esmpro23" in result.output
    assert "Template: generated automatically" in result.output
    assert "Validation: PASS" in result.output
    assert (work_dir / "NEC_ESMPRO23.xlsx").exists()
    assert output_path.exists()

    workbook = load_workbook(output_path, data_only=False)
    assert workbook.sheetnames == ["環境", "監視設定", "チェックシート"]
    assert workbook["環境"]["B8"].value == 1
    assert workbook["環境"]["B9"].value == 17
    assert workbook["チェックシート"]["C6"].value == "generated_template"
    monitor_types = {workbook["監視設定"][f"F{row}"].value for row in range(4, workbook["監視設定"].max_row + 1) if workbook["監視設定"][f"F{row}"].value not in (None, "")}
    assert "サービス監視" in monitor_types
    assert "イベントログ監視" in monitor_types
    critical_values = {workbook["監視設定"][f"J{row}"].value for row in range(4, workbook["監視設定"].max_row + 1) if workbook["監視設定"][f"J{row}"].value not in (None, "")}
    assert "5回連続注意" in critical_values
