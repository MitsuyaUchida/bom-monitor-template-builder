from pathlib import Path

from openpyxl import load_workbook  # type: ignore[import-untyped]

from bom_monitor_builder.cab_excel.cli import build_workbooks, parse_input_to_model


def test_build_workbook_from_extracted_fixture(tmp_path: Path) -> None:
    source = Path("tests/fixtures/sample_extracted")
    output = tmp_path / "sample.xlsx"

    output_paths = build_workbooks(source, output=output, extracted=True, overwrite=True)

    assert output_paths == [output]
    workbook = load_workbook(output)
    assert workbook.sheetnames == [
        "表紙",
        "監視グループ一覧",
        "監視項目一覧",
        "監視項目詳細",
        "XML全項目",
        "解析情報",
    ]
    assert workbook["表紙"]["B5"].value == "8.0"
    assert workbook["監視グループ一覧"]["C2"].value == "ディスクパフォーマンス監視"
    assert workbook["監視項目一覧"]["D2"].value == "レイテンシ(ディスク転送時間)"
    assert workbook["監視項目一覧"]["G2"].value == "5分"
    assert workbook["監視項目一覧"]["I2"].value == "20 以上"
    assert workbook["監視項目一覧"]["J2"].value == "30 以上"
    assert workbook["監視項目一覧"]["M2"].value == "powershell.exe"
    assert workbook["監視項目一覧"].freeze_panes == "A2"
    assert workbook["監視項目一覧"].auto_filter.ref is not None


def test_parse_actual_template_data_cab() -> None:
    input_path = Path("import/TemplateData/0002_Windows 基本/0103_ハードディスク負荷状況.cab")

    parsed, session = parse_input_to_model(input_path, extracted=False, keep_extracted=False)
    try:
        assert parsed.manifest.product == "BOM for Windows"
        assert parsed.manifest.export_type == "Monitor_Export"
        assert parsed.manifest.bom_version == "8.0"
        assert len(parsed.groups) == 1
        assert len(parsed.items) == 1
        assert parsed.items[0].raw_values["Type"] == "DiskQueueLength"
    finally:
        session.cleanup()
