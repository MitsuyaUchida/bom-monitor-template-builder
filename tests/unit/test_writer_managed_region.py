from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.comments import Comment  # type: ignore[import-untyped]
from openpyxl.styles import PatternFill  # type: ignore[import-untyped]

from bom_monitor_builder.design_excel.models import MonitorGroup, MonitorItem, RenderPlan, RenderRow, RenderTable, WorkbookModel
from bom_monitor_builder.design_excel.writer import write_render_plan


def test_writer_clears_stale_values_when_template_has_more_rows_than_model(tmp_path: Path) -> None:
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Settings"
    fill = PatternFill(fill_type="solid", fgColor="00D9EAF7")
    for row in range(3, 7):
        for column, value in {"A": "GXX", "B": "Template Group", "C": f"OLD{row}", "D": f"Old Monitor {row}"}.items():
            cell = sheet[f"{column}{row}"]
            cell.value = value
            cell.fill = fill
    sheet["D6"].comment = Comment("stale", "tester")
    sheet["D6"].hyperlink = "https://example.com/stale"
    workbook.save(template_path)

    plan = build_plan(
        template_path=template_path,
        output_path=output_path,
        rows=[
            {"A": "G01", "B": "Platform", "C": "MON01", "D": "CPU"},
            {"A": "G01", "B": "Platform", "C": "MON02", "D": "Memory"},
            {"A": "G02", "B": "Apps", "C": "MON03", "D": "Disk"},
        ],
        managed_columns={"group_id": "A", "group_name": "B", "monitor_id": "C", "monitor_name": "D"},
    )

    write_render_plan(plan)

    result = load_workbook(output_path)
    sheet = result["Settings"]
    assert sheet["C3"].value == "MON01"
    assert sheet["C5"].value == "MON03"
    assert sheet["C6"].value is None
    assert sheet["D6"].value is None
    assert sheet["D6"].comment is None
    assert sheet["D6"].hyperlink is None
    assert sheet["A5"].fill.fgColor.rgb == "00D9EAF7"


def test_writer_adds_rows_with_template_style_without_copying_stale_values(tmp_path: Path) -> None:
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Settings"
    fill = PatternFill(fill_type="solid", fgColor="00FFD966")
    for row in range(3, 5):
        for column, value in {"A": "GX", "B": "Template Group", "C": "OLD", "D": "Template Monitor"}.items():
            cell = sheet[f"{column}{row}"]
            cell.value = value
            cell.fill = fill
        sheet.row_dimensions[row].height = 24
    workbook.save(template_path)

    plan = build_plan(
        template_path=template_path,
        output_path=output_path,
        rows=[
            {"A": "G01", "B": "Platform", "C": "MON01", "D": "CPU"},
            {"A": "G01", "B": "Platform", "C": "MON02", "D": "Memory"},
            {"A": "G02", "B": "Apps", "C": "MON03", "D": "Disk"},
            {"A": "G02", "B": "Apps", "C": "MON04", "D": "Service"},
            {"A": "G03", "B": "DB", "C": "MON05", "D": "Event"},
        ],
        managed_columns={"group_id": "A", "group_name": "B", "monitor_id": "C", "monitor_name": "D"},
    )

    write_render_plan(plan)

    result = load_workbook(output_path)
    sheet = result["Settings"]
    assert sheet["C7"].value == "MON05"
    assert sheet["D7"].value == "Event"
    assert sheet["D7"].fill.fgColor.rgb == "00FFD966"
    assert sheet.row_dimensions[7].height == 24
    assert sheet["D7"].value != "Template Monitor"


def test_writer_rebuilds_group_merges_from_model(tmp_path: Path) -> None:
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Settings"
    for row in range(3, 7):
        sheet[f"A{row}"] = "OLD"
        sheet[f"B{row}"] = "OLD GROUP"
        sheet[f"C{row}"] = f"OLD{row}"
    sheet.merge_cells("A3:A6")
    sheet.merge_cells("B3:B6")
    workbook.save(template_path)

    plan = build_plan(
        template_path=template_path,
        output_path=output_path,
        rows=[
            {"A": "G01", "B": "Platform", "C": "MON01"},
            {"A": None, "B": None, "C": "MON02"},
            {"A": "G02", "B": "Apps", "C": "MON03"},
            {"A": None, "B": None, "C": "MON04"},
            {"A": None, "B": None, "C": "MON05"},
        ],
        managed_columns={"group_id": "A", "group_name": "B", "monitor_id": "C"},
        group_merge_fields=["group_id", "group_name"],
        group_sizes=[2, 3],
    )

    write_render_plan(plan)

    result = load_workbook(output_path)
    sheet = result["Settings"]
    merges = {str(item) for item in sheet.merged_cells.ranges}
    assert "A3:A4" in merges
    assert "A5:A7" in merges
    assert "B3:B4" in merges
    assert "B5:B7" in merges
    assert "A3:A6" not in merges


def build_plan(
    *,
    template_path: Path,
    output_path: Path,
    rows: list[dict[str, str | None]],
    managed_columns: dict[str, str],
    group_merge_fields: list[str] | None = None,
    group_sizes: list[int] | None = None,
) -> RenderPlan:
    group_sizes = group_sizes or [len(rows)]
    groups: list[MonitorGroup] = []
    items: list[MonitorItem] = []
    row_index = 0
    for group_number, group_size in enumerate(group_sizes, start=1):
        first_row = rows[row_index]
        group_id = str(first_row.get("A") or f"G{group_number:02d}")
        group_name = str(first_row.get("B") or f"Group {group_number}")
        groups.append(MonitorGroup(group_id=group_id, group_name=group_name, enabled=True, comment=None, source_row=group_number))
        for _ in range(group_size):
            row_values = rows[row_index]
            items.append(
                MonitorItem(
                    group_id=row_values.get("A") or group_id,
                    group_name=row_values.get("B") or group_name,
                    monitor_id=str(row_values.get("C")),
                    monitor_name=str(row_values.get("D", row_values.get("C"))),
                    monitor_type=None,
                    enabled="有効",
                    interval="5分",
                    warning_condition=None,
                    critical_condition=None,
                    comment=None,
                    source_sheet="Settings",
                    source_row=3 + row_index,
                )
            )
            row_index += 1
    model = WorkbookModel(
        source_path=Path("input/sample.xlsx"),
        source_name="sample.xlsx",
        groups=groups,
        items=items,
    )
    table_rows = [RenderRow(values=row, source_monitor_id=str(row.get("C"))) for row in rows]
    table = RenderTable(
        name="settings",
        sheet="Settings",
        start_row=3,
        template_row=3,
        rows=table_rows,
        managed_columns=managed_columns,
        group_merge_fields=group_merge_fields or [],
        layout={
            "group_merge": {
                "enabled": bool(group_merge_fields),
                "columns": [managed_columns[field_name] for field_name in (group_merge_fields or [])],
            }
        },
        clear_existing_data=True,
        trim_unused_rows=True,
        reuse_existing_rows=True,
    )
    return RenderPlan(
        output_path=output_path,
        template_path=template_path,
        cells={},
        tables=[table],
        model=model,
        profile={"security": {"sanitize_template_strings": False}},
    )
