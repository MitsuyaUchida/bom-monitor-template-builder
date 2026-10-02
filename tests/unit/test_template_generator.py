from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook  # type: ignore[import-untyped]

from bom_monitor_builder.design_excel.models import MonitorGroup, MonitorItem, WorkbookModel
from bom_monitor_builder.design_excel.template_generator import create_generated_render_plan


def test_generated_template_creates_expected_workbook_structure(tmp_path: Path) -> None:
    model = build_model()
    profile = {
        "profile": {"id": "nec_esmpro23", "name": "NEC ESMPRO23"},
        "template": {"resolved_path": None},
        "security": {"secret_patterns": ["password", "token"]},
        "output": {
            "generated_template": {
                "enabled": True,
                "columns": [
                    {"key": "no", "label": "No."},
                    {"key": "group_id", "label": "グループID"},
                    {"key": "group_name", "label": "グループ名"},
                    {"key": "monitor_id", "label": "監視ID"},
                    {"key": "monitor_name", "label": "監視名"},
                    {"key": "monitor_type", "label": "監視種別"},
                    {"key": "enabled", "label": "有効/無効"},
                    {"key": "interval", "label": "監視間隔"},
                    {"key": "warning_condition", "label": "注意条件"},
                    {"key": "critical_condition", "label": "危険条件"},
                    {"key": "target", "label": "監視対象"},
                    {"key": "remarks", "label": "備考"},
                ]
            }
        },
        "transform": {"defaults": {}, "replacements": {}, "normalizers": {}, "derived_fields": {}, "row_rules": []},
    }

    plan = create_generated_render_plan(model, profile, tmp_path / "output.xlsx")

    assert plan.template_source == "generated_template"
    assert plan.template_path is not None
    workbook = load_workbook(plan.template_path, data_only=False)
    assert workbook.sheetnames == ["環境", "監視設定", "チェックシート"]

    settings = workbook["監視設定"]
    assert settings["A3"].value == "No."
    assert settings["B3"].value == "グループID"
    assert settings["C3"].value == "グループ名"
    assert settings.freeze_panes == "A4"
    assert settings.column_dimensions["A"].width == 8
    assert settings.column_dimensions["C"].width == 28
    assert settings.row_dimensions[3].height == 24
    assert settings["A1"].fill.fill_type == "solid"
    assert settings["A1"].fill.fgColor.type == "rgb"
    assert settings["A1"].fill.fgColor.rgb == "FFFFFCCC"
    assert settings["A3"].fill.fgColor.rgb == "FFFFFCCC"
    assert settings["A4"].alignment.horizontal == "center"
    assert settings["C4"].alignment.horizontal == "left"
    assert settings["C4"].alignment.wrap_text is True
    assert settings["A4"].border.top.style == "thin"
    assert settings["L4"].border.bottom.style == "thin"


def test_generated_template_plan_hides_group_labels_after_first_row() -> None:
    model = build_model()
    profile = {
        "profile": {"id": "nec_esmpro23", "name": "NEC ESMPRO23"},
        "template": {"resolved_path": None},
        "security": {"secret_patterns": ["password"]},
        "output": {"generated_template": {"enabled": True}},
        "transform": {"defaults": {}, "replacements": {}, "normalizers": {}, "derived_fields": {}, "row_rules": []},
    }

    plan = create_generated_render_plan(model, profile, Path("output.xlsx"))

    assert len(plan.tables) == 1
    table = plan.tables[0]
    assert len(table.rows) == 3
    assert table.rows[0].values["B"] == "GRP01"
    assert table.rows[0].values["C"] == "グループA"
    assert table.rows[1].values["B"] == "GRP01"
    assert table.rows[2].values["B"] == "GRP02"
    assert table.group_labels_first_row_only is True


def build_model() -> WorkbookModel:
    return WorkbookModel(
        source_path=Path("input/nec.xlsx"),
        source_name="nec.xlsx",
        groups=[
            MonitorGroup(group_id="GRP01", group_name="グループA", enabled="有効", comment=None, source_row=2),
            MonitorGroup(group_id="GRP02", group_name="グループB", enabled="有効", comment=None, source_row=3),
        ],
        items=[
            MonitorItem(
                group_id="GRP01",
                group_name="グループA",
                monitor_id="MON01",
                monitor_name="サービス監視",
                monitor_type="Service",
                enabled="有効",
                interval="5分",
                warning_condition="停止",
                critical_condition="停止",
                comment=None,
                details={"ObjectName": "SvcA"},
            ),
            MonitorItem(
                group_id="GRP01",
                group_name="グループA",
                monitor_id="MON02",
                monitor_name="イベント監視",
                monitor_type="EventlogWSA",
                enabled="有効",
                interval="5分",
                warning_condition="1件以上",
                critical_condition="5件以上",
                comment=None,
                details={"ObjectName": "AppLog"},
            ),
            MonitorItem(
                group_id="GRP02",
                group_name="グループB",
                monitor_id="MON03",
                monitor_name="性能監視",
                monitor_type="Perf",
                enabled="無効",
                interval="10分",
                warning_condition="80%以上",
                critical_condition="90%以上",
                comment=None,
                details={"ObjectName": "CPU"},
            ),
        ],
    )
