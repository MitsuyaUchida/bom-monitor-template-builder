from __future__ import annotations

import pytest
from openpyxl import Workbook, load_workbook

from bom_monitor_builder.design_excel.condition_display import (
    get_condition_display_name,
    get_service_state_display_name,
)
from bom_monitor_builder.design_excel.mapper import build_table_row


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("1", "停止"),
        ("2", "開始中"),
        ("3", "停止中"),
        ("4", "開始"),
        ("5", "再開中"),
        ("6", "一時停止中"),
        ("7", "一時停止"),
        ("99", "99"),
    ],
)
def test_current_state_display(value: str, expected: str) -> None:
    assert get_service_state_display_name(
        value, monitor_type="Service", value_name="CurrentState"
    ) == expected


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("2ContYellow", "2回連続注意"),
        ("3ContYellow", "3回連続注意"),
        ("4ContYellow", "4回連続注意"),
        ("10ContYellow", "10回連続注意"),
    ],
)
def test_cont_yellow_display(value: str, expected: str) -> None:
    assert get_condition_display_name(value) == expected


def test_get_condition_display_name_for_unknown_value() -> None:
    assert get_condition_display_name("UnknownCondition") == "UnknownCondition"


def test_get_condition_display_name_for_empty_string() -> None:
    assert get_condition_display_name("") == ""


def test_get_condition_display_name_for_none() -> None:
    assert get_condition_display_name(None) == ""


def test_table_mapping_applies_display_conversion_without_mutating_input() -> None:
    values = {
        "monitor_type": "Service",
        "ValueName": "CurrentState",
        "warning_condition": "4",
        "critical_condition": "1",
    }

    row = build_table_row(
        values,
        {"H": "warning_condition", "I": "critical_condition"},
    )

    assert row == {"H": "開始", "I": "停止"}
    assert values["warning_condition"] == "4"
    assert values["critical_condition"] == "1"


def test_condition_mapping_preserves_unknown_current_state() -> None:
    assert get_service_state_display_name(
        "99", monitor_type="Service", value_name="CurrentState"
    ) == "99"


def test_condition_display_values_are_written_to_excel_cells(tmp_path) -> None:
    values = {
        "monitor_type": "Service",
        "ValueName": "CurrentState",
        "warning_condition": "4",
        "critical_condition": "5 ContYellow",
    }
    row = build_table_row(values, {"H": "warning_condition", "J": "critical_condition"})
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["H5"] = row["H"]
    worksheet["J5"] = row["J"]
    stopped_row = build_table_row(
        {"monitor_type": "Service", "ValueName": "CurrentState", "warning_condition": "1"},
        {"H": "warning_condition"},
    )
    worksheet["H6"] = stopped_row["H"]
    path = tmp_path / "display.xlsx"
    workbook.save(path)

    saved_sheet = load_workbook(path, data_only=True).active
    assert saved_sheet["H5"].value == "\u958b\u59cb"
    assert saved_sheet["H6"].value == "\u505c\u6b62"
    assert saved_sheet["J5"].value == "5\u56de\u9023\u7d9a\u6ce8\u610f"


def test_current_state_mapping_requires_service_context() -> None:
    assert get_service_state_display_name(
        "4", monitor_type="Perf", value_name="CurrentState"
    ) == "4"
    assert get_service_state_display_name(
        "4", monitor_type="Service", value_name="OtherValue"
    ) == "4"
