from pathlib import Path

from openpyxl import Workbook  # type: ignore[import-untyped]

from bom_monitor_builder.design_excel.parser import parse_workbook


def test_parse_workbook_uses_header_aliases(tmp_path: Path) -> None:
    workbook = Workbook()
    groups = workbook.active
    groups.title = "グループ一覧"
    groups.append(["番号", "グループID", "グループ", "Enabled", "Comment"])
    groups.append([1, "GRP10", "AliasGroup", "true", "group comment"])

    monitors = workbook.create_sheet("監視一覧")
    monitors.append(["No", "グループ", "MonitorID", "名称", "タイプ", "Enabled", "Interval", "Warning", "Critical", "Comment"])
    monitors.append([1, "AliasGroup", "MON10.xml", "CPU Monitor", "Perf", "false", "5 分", "80 以上", "90 以上", "item comment"])

    details = workbook.create_sheet("監視詳細")
    details.append(["監視No.", 1])
    details.append(["XMLファイルパス", "Monitor/GRP10/MON10.xml"])
    details.append(["Options", "-pw:secret"])

    input_path = tmp_path / "input.xlsx"
    workbook.save(input_path)

    profile = {
        "source": {
            "header_search_rows": 5,
            "sheets": {
                "groups": {"candidates": ["グループ一覧"]},
                "monitors": {"candidates": ["監視一覧"]},
                "details": {"candidates": ["監視詳細"], "required": False},
            },
            "fields": {
                "group_id": {"sections": ["groups"], "aliases": ["グループID"]},
                "group_name": {"sections": ["groups", "monitors"], "aliases": ["グループ名", "グループ"]},
                "enabled": {"sections": ["groups", "monitors"], "aliases": ["有効", "Enabled"]},
                "comment": {"sections": ["groups", "monitors"], "aliases": ["コメント", "Comment"]},
                "monitor_id": {"sections": ["monitors"], "aliases": ["監視ファイル", "MonitorID"]},
                "monitor_name": {"sections": ["monitors"], "aliases": ["監視名", "名称"]},
                "monitor_type": {"sections": ["monitors"], "aliases": ["監視タイプ", "タイプ"]},
                "interval": {"sections": ["monitors"], "aliases": ["監視間隔", "Interval"]},
                "warning_condition": {"sections": ["monitors"], "aliases": ["注意判定", "Warning"]},
                "critical_condition": {"sections": ["monitors"], "aliases": ["危険判定", "Critical"]},
            },
        }
    }

    model = parse_workbook(input_path, profile)

    assert len(model.groups) == 1
    assert model.groups[0].group_id == "GRP10"
    assert len(model.items) == 1
    assert model.items[0].monitor_id == "MON10.xml"
    assert model.items[0].details["Options"] == "-pw:secret"
