from __future__ import annotations

import json
from pathlib import Path

from click.testing import CliRunner
from openpyxl import load_workbook  # type: ignore[import-untyped]

from bom_monitor_builder.cli import main


def test_template_scan_writes_iteration00_outputs(tmp_path: Path) -> None:
    source_path = create_template_data(tmp_path)
    output_path = tmp_path / "knowledge"

    runner = CliRunner()
    result = runner.invoke(
        main,
        [
            "template",
            "scan",
            "--source",
            str(source_path),
            "--output",
            str(output_path),
        ],
    )

    assert result.exit_code == 0
    assert "TemplateData Scan" in result.output
    assert "Categories: 2" in result.output
    assert "CAB: 2" in result.output
    assert "HTML: 2" in result.output
    assert "YAML: 1" in result.output
    assert "Images: 1" in result.output
    assert "Other files: 1" in result.output

    expected_files = [
        output_path / "inventory/directory_tree.json",
        output_path / "inventory/files.json",
        output_path / "inventory/categories.json",
        output_path / "inventory/templates.json",
        output_path / "inventory/extensions.json",
        output_path / "inventory/summary.json",
        output_path / "reports/category_inventory.csv",
        output_path / "reports/template_inventory.csv",
        output_path / "reports/file_inventory.csv",
        output_path / "reports/template_summary.xlsx",
    ]
    for path in expected_files:
        assert path.exists()

    summary = json.loads((output_path / "inventory/summary.json").read_text(encoding="utf-8"))
    assert summary["category_count"] == 2
    assert summary["autotemplate"]["exists"] is True
    assert summary["autotemplate"]["yaml_loadable"] is True
    assert summary["autotemplate"]["encoding"] == "utf-8"

    categories = json.loads((output_path / "inventory/categories.json").read_text(encoding="utf-8"))
    assert categories["categories"] == [
        {
            "cab_count": 1,
            "category_id": "0001",
            "category_name": "Base",
            "directory": "0001_Base",
            "file_count": 3,
            "html_count": 1,
        },
        {
            "cab_count": 1,
            "category_id": "0002",
            "category_name": "App",
            "directory": "0002_App",
            "file_count": 3,
            "html_count": 1,
        },
    ]

    templates = json.loads((output_path / "inventory/templates.json").read_text(encoding="utf-8"))
    template_rows = {row["template_name_candidate"]: row for row in templates["templates"]}
    assert template_rows["server"]["same_name_pair"] is True
    assert template_rows["agent"]["html_only"] is True
    assert template_rows["db"]["cab_only"] is True

    workbook = load_workbook(output_path / "reports/template_summary.xlsx")
    assert workbook.sheetnames == ["Summary", "Categories", "Templates", "Files"]


def test_template_scan_marks_invalid_autotemplate_yaml(tmp_path: Path) -> None:
    source_path = create_template_data(tmp_path)
    (source_path / "AutoTemplate.yml").write_text("catalog: [1,\n", encoding="utf-8")
    output_path = tmp_path / "knowledge"

    runner = CliRunner()
    result = runner.invoke(
        main,
        [
            "template",
            "scan",
            "--source",
            str(source_path),
            "--output",
            str(output_path),
        ],
    )

    assert result.exit_code == 0
    summary = json.loads((output_path / "inventory/summary.json").read_text(encoding="utf-8"))
    assert summary["autotemplate"]["yaml_loadable"] is False
    assert summary["autotemplate"]["yaml_error"] is not None


def create_template_data(tmp_path: Path) -> Path:
    source_path = tmp_path / "import" / "TemplateData"
    base = source_path / "0001_Base"
    app = source_path / "0002_App"
    nested = app / "nested"
    base.mkdir(parents=True)
    nested.mkdir(parents=True)

    (source_path / "AutoTemplate.yml").write_text("catalog:\n  - name: test\n", encoding="utf-8")
    (base / "server.cab").write_bytes(b"cab-data")
    (base / "server.htm").write_text("<html>server</html>", encoding="utf-8")
    (base / "図.png").write_bytes(b"\x89PNG\r\n\x1a\n")
    (app / "agent.html").write_text("<html>agent</html>", encoding="utf-8")
    (nested / "db.cab").write_bytes(b"db-cab")
    (nested / "readme.txt").write_text("note", encoding="utf-8")
    return source_path
