from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]

from .exceptions import ParseWorkbookError
from .models import MonitorGroup, MonitorItem, WorkbookModel


def parse_workbook(input_path: Path, profile: dict[str, Any]) -> WorkbookModel:
    if not input_path.exists():
        raise ParseWorkbookError(f"Input workbook not found: {input_path}")
    workbook = load_workbook(input_path, data_only=False)
    source_config = profile["source"]
    group_sheet = find_sheet(workbook.sheetnames, source_config["sheets"].get("groups", {}))
    monitor_sheet = find_sheet(workbook.sheetnames, source_config["sheets"].get("monitors", {}))
    detail_sheet_config = source_config["sheets"].get("details", {})
    detail_sheet_name = find_sheet(workbook.sheetnames, detail_sheet_config, required=detail_sheet_config.get("required", False))

    groups = parse_tabular_sheet(workbook[group_sheet], source_config["fields"], source_config.get("header_search_rows", 10), "groups")
    items = parse_tabular_sheet(workbook[monitor_sheet], source_config["fields"], source_config.get("header_search_rows", 10), "monitors")
    detail_map: dict[str, dict[str, Any]] = {}
    if detail_sheet_name:
        detail_map = parse_detail_sheet(workbook[detail_sheet_name])
    actions: list[dict[str, Any]] = []
    if "Actions" in workbook.sheetnames:
        action_sheet = workbook["Actions"]
        headers = [str(cell.value or "") for cell in action_sheet[1]]
        actions = [
            {header: action_sheet.cell(row_idx, column_idx).value for column_idx, header in enumerate(headers, start=1) if header}
            for row_idx in range(2, action_sheet.max_row + 1)
            if any(action_sheet.cell(row_idx, column_idx).value is not None for column_idx in range(1, action_sheet.max_column + 1))
        ]
    profile_id = (
        profile.get("profile", {}).get("id")
        if isinstance(profile.get("profile"), dict)
        else None
    ) or "generic"
    group_id_by_name = {group.group_name: group.group_id for group in groups}
    for item in items:
        item.group_id = group_id_by_name.get(item.group_name or "")
        for detail_key in build_detail_lookup_keys(item):
            if detail_key in detail_map:
                item.details = detail_map[detail_key]
                break
    return WorkbookModel(
        source_path=input_path,
        source_name=input_path.name,
        groups=groups,
        items=items,
        details=detail_map,
        metadata={"sheetnames": workbook.sheetnames},
        extensions={
            profile_id: {
                "details": detail_map,
            }
        },
        raw_sections={"actions": actions},
    )


def parse_tabular_sheet(
    worksheet: Worksheet,
    fields: dict[str, Any],
    header_search_rows: int,
    section: str,
) -> list[Any]:
    header_row_index, field_columns = detect_header_row(worksheet, fields, header_search_rows, section)
    rows: list[Any] = []
    for row_idx in range(header_row_index + 1, worksheet.max_row + 1):
        row_values = {
            field_name: worksheet.cell(row_idx, column_index).value
            for field_name, column_index in field_columns.items()
        }
        if should_skip_row(row_values):
            continue
        raw_values = {
            normalize_header(str(worksheet.cell(header_row_index, col).value or f"column_{col}")): worksheet.cell(row_idx, col).value
            for col in range(1, worksheet.max_column + 1)
        }
        if section == "groups":
            rows.append(
                MonitorGroup(
                    group_id=str_or_empty(row_values.get("group_id")),
                    group_name=str_or_empty(row_values.get("group_name")),
                    enabled=row_values.get("enabled"),
                    comment=optional_str(row_values.get("comment")),
                    source_row=row_idx,
                    raw_values=raw_values,
                )
            )
        else:
            rows.append(
                MonitorItem(
                    group_id=None,
                    group_name=optional_str(row_values.get("group_name")),
                    monitor_id=str_or_empty(row_values.get("monitor_id")),
                    monitor_name=str_or_empty(row_values.get("monitor_name")),
                    monitor_type=optional_str(row_values.get("monitor_type")),
                    enabled=row_values.get("enabled"),
                    interval=optional_str(row_values.get("interval")),
                    warning_condition=optional_str(row_values.get("warning_condition")),
                    critical_condition=optional_str(row_values.get("critical_condition")),
                    comment=optional_str(row_values.get("comment")),
                    source_sheet=worksheet.title,
                    source_row=row_idx,
                    raw_values=raw_values,
                )
            )
    return rows


def parse_detail_sheet(worksheet: Worksheet) -> dict[str, dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    values: dict[str, Any] = {}
    for row_idx in range(1, worksheet.max_row + 1):
        key = worksheet.cell(row_idx, 1).value
        value = worksheet.cell(row_idx, 2).value
        if key is None:
            continue
        normalized_key = str(key).strip()
        if normalized_key == "監視No." and values:
            blocks.append(values)
            values = {}
        values[normalized_key] = value
    if values:
        blocks.append(values)
    detail_map: dict[str, dict[str, Any]] = {}
    for block in blocks:
        detail_keys: list[str] = []
        monitor_id: str | None = None
        group_id: str | None = None
        value = block.get("XMLファイルパス")
        if isinstance(value, str):
            match = re.search(r"/([^/]+)/([^/]+)\.xml$", value)
            if match:
                group_id = match.group(1)
                monitor_id = match.group(2)
        if not monitor_id:
            monitor_file = block.get("監視ファイル")
            if isinstance(monitor_file, str):
                monitor_id = re.sub(r"\.[A-Za-z0-9]+$", "", monitor_file)
        if not monitor_id:
            monitor_id = str(block.get("監視No.", f"detail_{len(detail_map)+1}"))
        if group_id:
            detail_keys.append(f"{group_id}/{monitor_id}")
        group_name = block.get("グループ名")
        if isinstance(group_name, str) and group_name:
            detail_keys.append(f"{group_name}/{monitor_id}")
        detail_keys.append(monitor_id)
        for detail_key in detail_keys:
            detail_map[detail_key] = block
    return detail_map


def build_detail_lookup_keys(item: MonitorItem) -> list[str]:
    raw_monitor_id = item.monitor_id
    normalized_monitor_id = re.sub(r"\.[A-Za-z0-9]+$", "", raw_monitor_id)
    keys: list[str] = []
    if item.group_id:
        keys.append(f"{item.group_id}/{normalized_monitor_id}")
    if item.group_name:
        keys.append(f"{item.group_name}/{normalized_monitor_id}")
    keys.append(normalized_monitor_id)
    if raw_monitor_id != normalized_monitor_id:
        if item.group_id:
            keys.append(f"{item.group_id}/{raw_monitor_id}")
        if item.group_name:
            keys.append(f"{item.group_name}/{raw_monitor_id}")
        keys.append(raw_monitor_id)
    return keys


def find_sheet(sheetnames: list[str], config: dict[str, Any], required: bool = True) -> str | None:
    candidates = config.get("candidates", [])
    for candidate in candidates:
        if candidate in sheetnames:
            return candidate
    if required:
        raise ParseWorkbookError(f"Required sheet not found. Candidates: {candidates}")
    return None


def detect_header_row(
    worksheet: Worksheet,
    fields: dict[str, Any],
    header_search_rows: int,
    section: str,
) -> tuple[int, dict[str, int]]:
    candidates: list[tuple[int, int, dict[str, int]]] = []
    for row_idx in range(1, min(worksheet.max_row, header_search_rows) + 1):
        matched: dict[str, int] = {}
        for col_idx in range(1, worksheet.max_column + 1):
            value = worksheet.cell(row_idx, col_idx).value
            normalized = normalize_header(str(value)) if value is not None else ""
            for field_name, config in fields.items():
                aliases = config.get("aliases", [])
                sections = config.get("sections")
                if sections and section not in sections:
                    continue
                if normalized and normalized in {normalize_header(alias) for alias in aliases}:
                    matched.setdefault(field_name, col_idx)
        candidates.append((len(matched), row_idx, matched))
    candidates.sort(key=lambda item: item[0], reverse=True)
    if not candidates or candidates[0][0] == 0:
        raise ParseWorkbookError(f"Header row not found in sheet: {worksheet.title}")
    top_score = candidates[0][0]
    ambiguous = [item for item in candidates if item[0] == top_score]
    if len(ambiguous) > 1:
        rows = [item[1] for item in ambiguous]
        raise ParseWorkbookError(f"Ambiguous header row in sheet {worksheet.title}: {rows}")
    _, row_idx, matched = candidates[0]
    return row_idx, matched


def normalize_header(value: str) -> str:
    return re.sub(r"[\s\u3000\r\n]+", "", value.strip())


def should_skip_row(values: dict[str, Any]) -> bool:
    non_empty = [value for value in values.values() if value not in (None, "")]
    if not non_empty:
        return True
    primary = next(iter(non_empty))
    if isinstance(primary, str) and primary.strip().startswith("合計"):
        return True
    return False


def str_or_empty(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def optional_str(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None
