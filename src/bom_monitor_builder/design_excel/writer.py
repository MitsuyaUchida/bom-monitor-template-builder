from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.cell.cell import MergedCell  # type: ignore[import-untyped]
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]

from .exceptions import TemplateWriteError
from .layout_engine import build_row_contexts, expected_merge_ranges
from .models import RenderPlan, RenderTable
from .security import DEFAULT_SECRET_PATTERNS, sanitize_workbook_strings
from .style_engine import apply_table_styles


def write_render_plan(plan: RenderPlan) -> Path:
    workbook = load_workbook(plan.template_path)
    for table in plan.tables:
        worksheet = workbook[table.sheet]
        write_table(worksheet, table, plan)
    for cell_config in plan.cells.values():
        worksheet = workbook[cell_config["sheet"]]
        worksheet[cell_config["cell"]] = cell_config["value"]
    if plan.profile.get("security", {}).get("sanitize_template_strings", True):
        sanitize_workbook_strings(
            workbook,
            list(plan.profile.get("security", {}).get("secret_patterns", DEFAULT_SECRET_PATTERNS)),
        )
    plan.output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(plan.output_path)
    return plan.output_path


def write_table(worksheet: Worksheet, table: RenderTable, plan: RenderPlan) -> None:
    contexts = build_row_contexts(plan.model, table)
    ensure_row_capacity(
        worksheet,
        table.start_row,
        table.template_row,
        len(table.rows),
        reuse_existing_rows=table.reuse_existing_rows,
    )
    if table.clear_existing_data:
        clear_managed_region(worksheet, table)
    context_by_row = {context.row: context for context in contexts}
    for offset, render_row in enumerate(table.rows):
        target_row = table.start_row + offset
        if target_row != table.template_row and not row_has_style_source(worksheet, target_row, table.managed_columns):
            clone_row(worksheet, table.template_row, target_row)
        context = context_by_row.get(target_row)
        if table.group_labels_first_row_only and context is not None and not context.is_group_first:
            group_id_column = table.managed_columns.get("group_id")
            group_name_column = table.managed_columns.get("group_name")
            if group_id_column is not None:
                render_row.values[group_id_column] = None
            if group_name_column is not None:
                render_row.values[group_name_column] = None
        for column, value in render_row.values.items():
            worksheet[f"{column}{target_row}"] = value
    apply_table_styles(worksheet, table, contexts)
    rebuild_group_merges(worksheet, table, contexts)
    if table.trim_unused_rows:
        clear_unused_rows(worksheet, table, len(table.rows))


def clear_managed_region(worksheet: Worksheet, table: RenderTable) -> None:
    managed_column_indexes = sorted(column_index(column) for column in table.managed_columns.values())
    unmerge_managed_ranges(worksheet, table, managed_column_indexes)
    for row in range(table.start_row, worksheet.max_row + 1):
        for column_name in table.managed_columns.values():
            clear_cell(worksheet[f"{column_name}{row}"])


def clear_unused_rows(worksheet: Worksheet, table: RenderTable, row_count: int) -> None:
    start_row = table.start_row + row_count
    end_row = table.template_data_end_row or worksheet.max_row
    if start_row > end_row:
        return
    if table.unused_rows_mode == "delete_rows":
        worksheet.delete_rows(start_row, end_row - start_row + 1)
        return
    for row in range(start_row, end_row + 1):
        for column_name in table.managed_columns.values():
            clear_cell(worksheet[f"{column_name}{row}"])


def rebuild_group_merges(worksheet: Worksheet, table: RenderTable, contexts: list[object]) -> None:
    group_merge = table.layout.get("group_merge", {})
    columns = list(group_merge.get("columns", [])) if isinstance(group_merge, dict) else []
    for column_name in columns:
        unmerge_column_ranges(worksheet, column_name, table.start_row)
    for ranges in expected_merge_ranges(table, contexts).values():
        for cell_range in ranges:
            worksheet.merge_cells(cell_range)


def merge_if_needed(worksheet: Worksheet, column_name: str, start_row: int, end_row: int) -> None:
    if end_row <= start_row:
        return
    worksheet.merge_cells(f"{column_name}{start_row}:{column_name}{end_row}")


def unmerge_managed_ranges(worksheet: Worksheet, table: RenderTable, managed_column_indexes: list[int]) -> None:
    for merged_range in list(worksheet.merged_cells.ranges):
        if merged_range.max_row < table.start_row:
            continue
        if merged_range.min_col > managed_column_indexes[-1] or merged_range.max_col < managed_column_indexes[0]:
            continue
        if any(merged_range.min_col <= index <= merged_range.max_col for index in managed_column_indexes):
            worksheet.unmerge_cells(str(merged_range))


def unmerge_column_ranges(worksheet: Worksheet, column_name: str, start_row: int) -> None:
    target_index = column_index(column_name)
    for merged_range in list(worksheet.merged_cells.ranges):
        if merged_range.max_row < start_row:
            continue
        if merged_range.min_col <= target_index <= merged_range.max_col:
            worksheet.unmerge_cells(str(merged_range))


def clear_cell(cell: object) -> None:
    if isinstance(cell, MergedCell):
        return
    cell.value = None
    cell.comment = None
    cell.hyperlink = None


def row_has_style_source(worksheet: Worksheet, row: int, managed_columns: dict[str, str]) -> bool:
    if not managed_columns:
        return False
    for column_name in managed_columns.values():
        cell = worksheet[f"{column_name}{row}"]
        if not isinstance(cell, MergedCell) and cell.style_id != 0:
            return True
    return False


def column_index(column_name: str) -> int:
    index = 0
    for char in column_name.upper():
        index = index * 26 + (ord(char) - ord("A") + 1)
    return index


def ensure_row_capacity(
    worksheet: Worksheet,
    start_row: int,
    template_row: int,
    required_count: int,
    *,
    reuse_existing_rows: bool = False,
) -> None:
    if required_count <= 1:
        return
    if reuse_existing_rows and has_existing_capacity(worksheet, start_row, required_count):
        return
    worksheet.insert_rows(start_row + 1, required_count - 1)
    move_merged_ranges_for_insert(worksheet, start_row + 1, required_count - 1)
    for row_index in range(start_row + 1, start_row + required_count):
        clone_row(worksheet, template_row, row_index)


def has_existing_capacity(worksheet: Worksheet, start_row: int, required_count: int) -> bool:
    last_row = start_row + required_count - 1
    return worksheet.max_row >= last_row


def clone_row(worksheet: Worksheet, source_row: int, target_row: int) -> None:
    source_dim = worksheet.row_dimensions[source_row]
    target_dim = worksheet.row_dimensions[target_row]
    target_dim.height = source_dim.height
    target_dim.hidden = source_dim.hidden
    for col_idx in range(1, worksheet.max_column + 1):
        source_cell = worksheet.cell(source_row, col_idx)
        target_cell = worksheet.cell(target_row, col_idx)
        if source_cell.data_type == "f":
            target_cell.value = source_cell.value
        else:
            target_cell.value = source_cell.value
        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)  # noqa: SLF001
        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format
        if source_cell.font:
            target_cell.font = copy(source_cell.font)
        if source_cell.fill:
            target_cell.fill = copy(source_cell.fill)
        if source_cell.border:
            target_cell.border = copy(source_cell.border)
        if source_cell.alignment:
            target_cell.alignment = copy(source_cell.alignment)
        if source_cell.protection:
            target_cell.protection = copy(source_cell.protection)
    copy_row_merged_ranges(worksheet, source_row, target_row)


def copy_row_merged_ranges(worksheet: Worksheet, source_row: int, target_row: int) -> None:
    row_offset = target_row - source_row
    ranges_to_add: list[str] = []
    for merged in list(worksheet.merged_cells.ranges):
        if merged.min_row == source_row and merged.max_row == source_row:
            ranges_to_add.append(
                f"{worksheet.cell(target_row, merged.min_col).coordinate}:{worksheet.cell(target_row, merged.max_col).coordinate}"
            )
        elif merged.min_row <= source_row <= merged.max_row and merged.max_row > merged.min_row:
            continue
    for merged_range in ranges_to_add:
        if merged_range not in {str(item) for item in worksheet.merged_cells.ranges}:
            worksheet.merge_cells(merged_range)


def move_merged_ranges_for_insert(worksheet: Worksheet, start_row: int, row_count: int) -> None:
    existing_ranges = list(worksheet.merged_cells.ranges)
    if not existing_ranges:
        return
    worksheet.merged_cells.ranges = []
    for merged in existing_ranges:
        if merged.min_row >= start_row:
            merged.shift(row_shift=row_count, col_shift=0)
        worksheet.merged_cells.add(merged)
