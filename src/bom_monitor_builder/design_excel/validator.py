from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]

from .exceptions import ValidationError
from .layout_engine import build_row_contexts, expected_merge_ranges
from .models import RenderPlan, ValidationReport
from .security import DEFAULT_SECRET_PATTERNS
from .style_engine import expected_alignment, expected_border, expected_row_height


def validate_output(
    output_path: Path,
    template_path: Path,
    input_path: Path,
    plan: RenderPlan,
) -> ValidationReport:
    output_wb = load_workbook(output_path, data_only=False)
    template_wb = load_workbook(template_path, data_only=False)
    input_wb = load_workbook(input_path, data_only=False)
    checks: list[str] = []

    if output_wb.sheetnames != template_wb.sheetnames:
        raise ValidationError("Sheet names or order differ from template.")
    checks.append("sheet order")
    for output_ws, template_ws in zip(output_wb.worksheets, template_wb.worksheets):
        if output_ws.sheet_state != template_ws.sheet_state:
            raise ValidationError(f"Sheet visibility changed: {output_ws.title}")
    checks.append("sheet visibility")
    for table in plan.tables:
        template_ws = template_wb[table.sheet]
        output_ws = output_wb[table.sheet]
        compare_sheet_structure(template_ws, output_ws, plan, table.start_row, len(table.rows))
    checks.append("template formatting")
    if not input_wb.sheetnames:
        raise ValidationError("Input workbook unreadable.")
    checks.append("input workbook readable")
    verify_written_values(output_wb, plan)
    checks.append("mapped values")
    verify_monitor_counts(output_wb, plan)
    checks.append("monitor count")
    verify_group_values(output_wb, plan)
    checks.append("group count")
    verify_no_stale_monitor_rows(output_wb, plan)
    checks.append("no stale monitor rows")
    verify_no_stale_template_data(output_wb, plan)
    checks.append("no stale template data")
    verify_merge_ranges(output_wb, plan)
    checks.append("merge ranges")
    verify_alignment_rules(output_wb, template_wb, plan)
    checks.append("alignment rules")
    verify_border_rules(output_wb, template_wb, plan)
    checks.append("border rules")
    verify_row_height_rules(output_wb, template_wb, plan)
    checks.append("row height rules")
    verify_no_secrets(
        output_wb,
        list(plan.profile.get("security", {}).get("secret_patterns", DEFAULT_SECRET_PATTERNS)),
        set(plan.profile.get("security", {}).get("allow_plaintext_cells", [])),
    )
    checks.append("secret masking")
    return ValidationReport(output_path=output_path, checks=checks)


def compare_sheet_structure(template_ws: Any, output_ws: Any, plan: RenderPlan, start_row: int, row_count: int) -> None:
    if output_ws.freeze_panes != template_ws.freeze_panes:
        raise ValidationError(f"Freeze panes changed in {output_ws.title}")
    if set(str(item) for item in template_ws.merged_cells.ranges if item.max_row < start_row or item.min_row > start_row) - set(
        str(item) for item in output_ws.merged_cells.ranges
    ):
        raise ValidationError(f"Merged cells were lost in {output_ws.title}")
    for column_key, dimension in template_ws.column_dimensions.items():
        if output_ws.column_dimensions[column_key].width != dimension.width:
            raise ValidationError(f"Column width changed in {output_ws.title}:{column_key}")
    for row_idx, dimension in template_ws.row_dimensions.items():
        if row_idx < start_row or row_idx >= start_row + max(row_count, 1):
            if output_ws.row_dimensions[row_idx].height != dimension.height:
                raise ValidationError(f"Row height changed in {output_ws.title}:{row_idx}")


def verify_written_values(workbook: Any, plan: RenderPlan) -> None:
    enforce_style_id_match = plan.template_source != "generated_template"
    for table in plan.tables:
        worksheet = workbook[table.sheet]
        if len(table.rows) != len(plan.model.items):
            raise ValidationError("Input/output monitor count mismatch.")
        for offset, row in enumerate(table.rows):
            target_row = table.start_row + offset
            reference_style_row = target_row if table.reuse_existing_rows else table.template_row
            for column, expected in row.values.items():
                actual = worksheet[f"{column}{target_row}"].value
                if actual != expected:
                    raise ValidationError(
                        f"Value mismatch at {table.sheet}!{column}{target_row}: expected={expected!r}, actual={actual!r}"
                    )
                if (
                    enforce_style_id_match
                    and worksheet[f"{column}{target_row}"].style_id != workbook[table.sheet][f"{column}{reference_style_row}"].style_id
                ):
                    raise ValidationError(f"Style changed at {table.sheet}!{column}{target_row}")
    for cell_config in plan.cells.values():
        worksheet = workbook[cell_config["sheet"]]
        if worksheet[cell_config["cell"]].value != cell_config["value"]:
            raise ValidationError(f"Cell mismatch at {cell_config['sheet']}!{cell_config['cell']}")


def verify_monitor_counts(workbook: Any, plan: RenderPlan) -> None:
    for table in plan.tables:
        monitor_column = table.managed_columns.get("monitor_id")
        if monitor_column is None:
            continue
        worksheet = workbook[table.sheet]
        actual_count = 0
        for row in range(table.start_row, worksheet.max_row + 1):
            if worksheet[f"{monitor_column}{row}"].value not in (None, ""):
                actual_count += 1
        if actual_count != len(plan.model.items):
            raise ValidationError(f"Monitor count mismatch in {table.sheet}: expected={len(plan.model.items)}, actual={actual_count}")


def verify_group_values(workbook: Any, plan: RenderPlan) -> None:
    for table in plan.tables:
        group_id_column = table.managed_columns.get("group_id")
        group_name_column = table.managed_columns.get("group_name")
        if group_id_column is None or group_name_column is None:
            continue
        worksheet = workbook[table.sheet]
        seen_groups: list[tuple[str, str]] = []
        for offset, row in enumerate(table.rows):
            row = table.start_row + offset
            group_id = worksheet[f"{group_id_column}{row}"].value
            group_name = worksheet[f"{group_name_column}{row}"].value
            if group_id in (None, "") and group_name in (None, ""):
                continue
            pair = (str(group_id), str(group_name))
            if not seen_groups or seen_groups[-1] != pair:
                seen_groups.append(pair)
        expected_groups: list[tuple[str, str]] = []
        for render_row in table.rows:
            expected_group_id = render_row.values.get(group_id_column)
            expected_group_name = render_row.values.get(group_name_column)
            if expected_group_id in (None, "") and expected_group_name in (None, ""):
                continue
            pair = (str(expected_group_id), str(expected_group_name))
            if not expected_groups or expected_groups[-1] != pair:
                expected_groups.append(pair)
        if seen_groups != expected_groups:
            raise ValidationError(f"Group sequence mismatch in {table.sheet}: expected={expected_groups!r}, actual={seen_groups!r}")


def verify_no_stale_monitor_rows(workbook: Any, plan: RenderPlan) -> None:
    for table in plan.tables:
        worksheet = workbook[table.sheet]
        stale_start = table.start_row + len(table.rows)
        for row in range(stale_start, worksheet.max_row + 1):
            for field_name, column_name in table.managed_columns.items():
                if worksheet[f"{column_name}{row}"].value not in (None, ""):
                    raise ValidationError(f"Stale monitor row data detected at {table.sheet}!{column_name}{row} ({field_name})")


def verify_no_stale_template_data(workbook: Any, plan: RenderPlan) -> None:
    for table in plan.tables:
        worksheet = workbook[table.sheet]
        expected_values = {
            (table.start_row + offset, column): value
            for offset, row in enumerate(table.rows)
            for column, value in row.values.items()
        }
        for (row, column), expected in expected_values.items():
            actual = worksheet[f"{column}{row}"].value
            if actual != expected:
                raise ValidationError(f"Stale template data detected at {table.sheet}!{column}{row}: expected={expected!r}, actual={actual!r}")


def verify_merge_ranges(workbook: Any, plan: RenderPlan) -> None:
    for table in plan.tables:
        contexts = build_row_contexts(plan.model, table)
        expected = {
            cell_range
            for ranges in expected_merge_ranges(table, contexts).values()
            for cell_range in ranges
        }
        actual = {
            str(item)
            for item in workbook[table.sheet].merged_cells.ranges
            if item.min_row >= table.start_row
        }
        if actual != expected:
            raise ValidationError(f"Merge range mismatch in {table.sheet}: expected={sorted(expected)!r}, actual={sorted(actual)!r}")


def verify_alignment_rules(workbook: Any, template_wb: Any, plan: RenderPlan) -> None:
    for table in plan.tables:
        contexts = build_row_contexts(plan.model, table)
        worksheet = workbook[table.sheet]
        template_ws = template_wb[table.sheet]
        for context in contexts:
            for column_name in table.managed_columns.values():
                actual = worksheet[f"{column_name}{context.row}"].alignment
                expected = expected_alignment(table, template_ws, context, column_name)
                if expected is None:
                    continue
                if (
                    actual.horizontal != expected.horizontal
                    or actual.vertical != expected.vertical
                    or actual.wrap_text != expected.wrap_text
                ):
                    raise ValidationError(f"Alignment mismatch at {table.sheet}!{column_name}{context.row}")


def verify_border_rules(workbook: Any, template_wb: Any, plan: RenderPlan) -> None:
    for table in plan.tables:
        contexts = build_row_contexts(plan.model, table)
        worksheet = workbook[table.sheet]
        template_ws = template_wb[table.sheet]
        for context in contexts:
            for column_name in table.managed_columns.values():
                actual = worksheet[f"{column_name}{context.row}"].border
                expected = expected_border(table, template_ws, context, column_name)
                if expected is None:
                    continue
                if (
                    actual.left.style != expected.left.style
                    or actual.right.style != expected.right.style
                    or actual.top.style != expected.top.style
                    or actual.bottom.style != expected.bottom.style
                ):
                    raise ValidationError(f"Border mismatch at {table.sheet}!{column_name}{context.row}")


def verify_row_height_rules(workbook: Any, template_wb: Any, plan: RenderPlan) -> None:
    for table in plan.tables:
        contexts = build_row_contexts(plan.model, table)
        worksheet = workbook[table.sheet]
        template_ws = template_wb[table.sheet]
        expected_height_ws = worksheet if plan.template_source == "generated_template" else template_ws
        for context in contexts:
            expected = expected_row_height(table, expected_height_ws, context)
            if expected is None:
                continue
            actual = worksheet.row_dimensions[context.row].height
            if actual != expected:
                raise ValidationError(f"Row height mismatch at {table.sheet}!{context.row}: expected={expected!r}, actual={actual!r}")


def verify_no_secrets(workbook: Any, patterns: list[str], allowed_cells: set[str]) -> None:
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                if f"{worksheet.title}!{cell.coordinate}" in allowed_cells:
                    continue
                if looks_like_unmasked_secret(cell.value, patterns):
                    raise ValidationError(f"Unmasked secret-like text detected at {worksheet.title}!{cell.coordinate}")


def looks_like_unmasked_secret(value: str, patterns: list[str]) -> bool:
    if "****" in value:
        return False
    if "-pw:" in value.lower():
        return True
    for pattern in patterns:
        if pattern.lower() not in value.lower():
            continue
        if any(token in value for token in ["=", ":"]):
            return True
    return False
