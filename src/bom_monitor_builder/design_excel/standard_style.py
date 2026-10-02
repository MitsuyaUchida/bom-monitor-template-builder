from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # type: ignore[import-untyped]

STANDARD_COLORS = {
    "title": "1F4E78",
    "section": "5B9BD5",
    "header": "D9EAF7",
    "group": "FFF2CC",
    "monitor_setting_background": "FFFFFCCC",
    "warning": "FFF2CC",
    "danger": "FCE4D6",
    "disabled": "E7E6E6",
    "white": "FFFFFF",
    "text": "000000",
}

STANDARD_FONT_NAME = "メイリオ"
STANDARD_FONT_SIZE = 11
STANDARD_THIN_SIDE = Side(style="thin", color=STANDARD_COLORS["text"])
STANDARD_MEDIUM_SIDE = Side(style="medium", color=STANDARD_COLORS["text"])


@dataclass(frozen=True, slots=True)
class CellStyle:
    font: Font
    fill: PatternFill
    alignment: Alignment
    border: Border


def title_style() -> CellStyle:
    return CellStyle(
        font=Font(name=STANDARD_FONT_NAME, size=14, bold=True, color=STANDARD_COLORS["white"]),
        fill=PatternFill("solid", fgColor=STANDARD_COLORS["title"]),
        alignment=Alignment(horizontal="left", vertical="center", wrap_text=True),
        border=Border(bottom=STANDARD_MEDIUM_SIDE),
    )


def monitor_setting_title_style() -> CellStyle:
    return CellStyle(
        font=Font(name=STANDARD_FONT_NAME, size=14, bold=True, color=STANDARD_COLORS["text"]),
        fill=PatternFill("solid", fgColor=STANDARD_COLORS["monitor_setting_background"]),
        alignment=Alignment(horizontal="left", vertical="center", wrap_text=True),
        border=Border(bottom=STANDARD_MEDIUM_SIDE),
    )


def section_style() -> CellStyle:
    return CellStyle(
        font=Font(name=STANDARD_FONT_NAME, size=11, bold=True, color=STANDARD_COLORS["white"]),
        fill=PatternFill("solid", fgColor=STANDARD_COLORS["section"]),
        alignment=Alignment(horizontal="left", vertical="center", wrap_text=True),
        border=Border(bottom=STANDARD_THIN_SIDE),
    )


def monitor_setting_section_style() -> CellStyle:
    return CellStyle(
        font=Font(name=STANDARD_FONT_NAME, size=11, bold=True, color=STANDARD_COLORS["text"]),
        fill=PatternFill("solid", fgColor=STANDARD_COLORS["monitor_setting_background"]),
        alignment=Alignment(horizontal="left", vertical="center", wrap_text=True),
        border=Border(bottom=STANDARD_THIN_SIDE),
    )


def header_style() -> CellStyle:
    return CellStyle(
        font=Font(name=STANDARD_FONT_NAME, size=10, bold=True, color=STANDARD_COLORS["text"]),
        fill=PatternFill("solid", fgColor=STANDARD_COLORS["header"]),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
        border=Border(
            left=STANDARD_THIN_SIDE,
            right=STANDARD_THIN_SIDE,
            top=STANDARD_THIN_SIDE,
            bottom=STANDARD_THIN_SIDE,
        ),
    )


def monitor_setting_header_style() -> CellStyle:
    return CellStyle(
        font=Font(name=STANDARD_FONT_NAME, size=10, bold=True, color=STANDARD_COLORS["text"]),
        fill=PatternFill("solid", fgColor=STANDARD_COLORS["monitor_setting_background"]),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
        border=Border(
            left=STANDARD_THIN_SIDE,
            right=STANDARD_THIN_SIDE,
            top=STANDARD_THIN_SIDE,
            bottom=STANDARD_THIN_SIDE,
        ),
    )


def group_style(horizontal: str = "left", wrap_text: bool = True) -> CellStyle:
    return CellStyle(
        font=Font(name=STANDARD_FONT_NAME, size=STANDARD_FONT_SIZE, bold=False, color=STANDARD_COLORS["text"]),
        fill=PatternFill("solid", fgColor=STANDARD_COLORS["group"]),
        alignment=Alignment(horizontal=horizontal, vertical="top", wrap_text=wrap_text),
        border=Border(
            left=STANDARD_THIN_SIDE,
            right=STANDARD_THIN_SIDE,
            top=STANDARD_MEDIUM_SIDE,
            bottom=STANDARD_THIN_SIDE,
        ),
    )


def monitor_setting_group_style(horizontal: str = "left", wrap_text: bool = True) -> CellStyle:
    return CellStyle(
        font=Font(name=STANDARD_FONT_NAME, size=STANDARD_FONT_SIZE, bold=False, color=STANDARD_COLORS["text"]),
        fill=PatternFill("solid", fgColor=STANDARD_COLORS["monitor_setting_background"]),
        alignment=Alignment(horizontal=horizontal, vertical="top", wrap_text=wrap_text),
        border=Border(
            left=STANDARD_THIN_SIDE,
            right=STANDARD_THIN_SIDE,
            top=STANDARD_MEDIUM_SIDE,
            bottom=STANDARD_THIN_SIDE,
        ),
    )


def warning_style(base_alignment: Alignment) -> CellStyle:
    return CellStyle(
        font=Font(name=STANDARD_FONT_NAME, size=STANDARD_FONT_SIZE, color=STANDARD_COLORS["text"]),
        fill=PatternFill("solid", fgColor=STANDARD_COLORS["warning"]),
        alignment=copy_alignment(base_alignment),
        border=Border(
            left=STANDARD_THIN_SIDE,
            right=STANDARD_THIN_SIDE,
            top=STANDARD_THIN_SIDE,
            bottom=STANDARD_THIN_SIDE,
        ),
    )


def monitor_setting_warning_style(base_alignment: Alignment) -> CellStyle:
    return CellStyle(
        font=Font(name=STANDARD_FONT_NAME, size=STANDARD_FONT_SIZE, color=STANDARD_COLORS["text"]),
        fill=PatternFill("solid", fgColor=STANDARD_COLORS["monitor_setting_background"]),
        alignment=copy_alignment(base_alignment),
        border=Border(
            left=STANDARD_THIN_SIDE,
            right=STANDARD_THIN_SIDE,
            top=STANDARD_THIN_SIDE,
            bottom=STANDARD_THIN_SIDE,
        ),
    )


def danger_style(base_alignment: Alignment) -> CellStyle:
    return CellStyle(
        font=Font(name=STANDARD_FONT_NAME, size=STANDARD_FONT_SIZE, color=STANDARD_COLORS["text"]),
        fill=PatternFill("solid", fgColor=STANDARD_COLORS["danger"]),
        alignment=copy_alignment(base_alignment),
        border=Border(
            left=STANDARD_THIN_SIDE,
            right=STANDARD_THIN_SIDE,
            top=STANDARD_THIN_SIDE,
            bottom=STANDARD_THIN_SIDE,
        ),
    )


def monitor_setting_danger_style(base_alignment: Alignment) -> CellStyle:
    return CellStyle(
        font=Font(name=STANDARD_FONT_NAME, size=STANDARD_FONT_SIZE, color=STANDARD_COLORS["text"]),
        fill=PatternFill("solid", fgColor=STANDARD_COLORS["monitor_setting_background"]),
        alignment=copy_alignment(base_alignment),
        border=Border(
            left=STANDARD_THIN_SIDE,
            right=STANDARD_THIN_SIDE,
            top=STANDARD_THIN_SIDE,
            bottom=STANDARD_THIN_SIDE,
        ),
    )


def disabled_style(base_alignment: Alignment, *, top_style: str | None = "thin", bottom_style: str | None = "thin") -> CellStyle:
    return CellStyle(
        font=Font(name=STANDARD_FONT_NAME, size=STANDARD_FONT_SIZE, color="666666"),
        fill=PatternFill("solid", fgColor=STANDARD_COLORS["disabled"]),
        alignment=copy_alignment(base_alignment),
        border=Border(
            left=STANDARD_THIN_SIDE,
            right=STANDARD_THIN_SIDE,
            top=copy_side(STANDARD_THIN_SIDE, top_style),
            bottom=copy_side(STANDARD_THIN_SIDE, bottom_style),
        ),
    )


def monitor_setting_disabled_style(
    base_alignment: Alignment,
    *,
    top_style: str | None = "thin",
    bottom_style: str | None = "thin",
) -> CellStyle:
    return CellStyle(
        font=Font(name=STANDARD_FONT_NAME, size=STANDARD_FONT_SIZE, color="666666"),
        fill=PatternFill("solid", fgColor=STANDARD_COLORS["monitor_setting_background"]),
        alignment=copy_alignment(base_alignment),
        border=Border(
            left=STANDARD_THIN_SIDE,
            right=STANDARD_THIN_SIDE,
            top=copy_side(STANDARD_THIN_SIDE, top_style),
            bottom=copy_side(STANDARD_THIN_SIDE, bottom_style),
        ),
    )


def data_style(
    *,
    horizontal: str,
    vertical: str,
    wrap_text: bool,
    top_style: str | None = "thin",
    bottom_style: str | None = "thin",
) -> CellStyle:
    return CellStyle(
        font=Font(name=STANDARD_FONT_NAME, size=STANDARD_FONT_SIZE, color=STANDARD_COLORS["text"]),
        fill=PatternFill("solid", fgColor=STANDARD_COLORS["white"]),
        alignment=Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap_text),
        border=Border(
            left=STANDARD_THIN_SIDE,
            right=STANDARD_THIN_SIDE,
            top=copy_side(STANDARD_THIN_SIDE, top_style),
            bottom=copy_side(STANDARD_THIN_SIDE, bottom_style),
        ),
    )


def monitor_setting_data_style(
    *,
    horizontal: str,
    vertical: str,
    wrap_text: bool,
    top_style: str | None = "thin",
    bottom_style: str | None = "thin",
) -> CellStyle:
    return CellStyle(
        font=Font(name=STANDARD_FONT_NAME, size=STANDARD_FONT_SIZE, color=STANDARD_COLORS["text"]),
        fill=PatternFill("solid", fgColor=STANDARD_COLORS["monitor_setting_background"]),
        alignment=Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap_text),
        border=Border(
            left=STANDARD_THIN_SIDE,
            right=STANDARD_THIN_SIDE,
            top=copy_side(STANDARD_THIN_SIDE, top_style),
            bottom=copy_side(STANDARD_THIN_SIDE, bottom_style),
        ),
    )


def apply_cell_style(cell: Any, style: CellStyle) -> None:
    cell.font = copy(style.font)
    cell.fill = copy(style.fill)
    cell.alignment = copy(style.alignment)
    cell.border = copy(style.border)


def copy_alignment(alignment: Alignment) -> Alignment:
    return Alignment(
        horizontal=alignment.horizontal,
        vertical=alignment.vertical,
        text_rotation=alignment.text_rotation,
        wrap_text=alignment.wrap_text,
        shrink_to_fit=alignment.shrink_to_fit,
        indent=alignment.indent,
    )


def copy_side(side: Side, style: str | None) -> Side:
    copied = copy(side)
    copied.style = style
    return copied
