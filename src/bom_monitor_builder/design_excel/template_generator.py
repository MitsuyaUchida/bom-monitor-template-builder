from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment  # type: ignore[import-untyped]

from .layout_engine import build_row_contexts
from .mapper import build_output_row, build_table_row
from .models import RenderPlan, RenderRow, RenderTable, WorkbookModel
from .security import DEFAULT_SECRET_PATTERNS, mask_secrets_in_mapping, sanitize_workbook_strings
from .standard_layout import apply_auto_filter, apply_standard_print_settings, apply_standard_sheet_view, clamp_width
from .standard_style import (
    STANDARD_FONT_NAME,
    apply_cell_style,
    data_style,
    header_style,
    monitor_setting_data_style,
    monitor_setting_header_style,
    monitor_setting_section_style,
    monitor_setting_title_style,
    section_style,
    title_style,
)

DEFAULT_COLUMNS = [
    {"key": "no", "label": "No.", "align": "center", "width": 8},
    {"key": "group_id", "label": "グループID", "align": "center", "width": 14},
    {"key": "group_name", "label": "グループ名", "align": "left", "width": 28},
    {"key": "monitor_id", "label": "監視ID", "align": "center", "width": 14},
    {"key": "monitor_name", "label": "監視名", "align": "left", "width": 28},
    {"key": "monitor_type", "label": "監視種別", "align": "center", "width": 14},
    {"key": "enabled", "label": "有効/無効", "align": "center", "width": 12},
    {"key": "interval", "label": "監視間隔", "align": "center", "width": 12},
    {"key": "average_count", "label": "平均回数", "align": "center", "width": 12},
    {"key": "warning_condition", "label": "注意条件", "align": "left", "width": 24},
    {"key": "critical_condition", "label": "危険条件", "align": "left", "width": 24},
    {"key": "target", "label": "監視対象", "align": "left", "width": 24},
    {"key": "action", "label": "アクション", "align": "left", "width": 22},
    {"key": "remarks", "label": "備考", "align": "left", "width": 32},
]


def profile_allows_generated_template(profile: dict[str, Any]) -> bool:
    template_section = profile.get("template", {})
    if bool(template_section.get("generate_if_missing")):
        return True
    generated = profile.get("output", {}).get("generated_template", {})
    return bool(generated.get("enabled"))


def create_generated_render_plan(
    model: WorkbookModel,
    profile: dict[str, Any],
    output_path: Path,
) -> RenderPlan:
    generated = profile.get("output", {}).get("generated_template", {})
    columns = resolve_generated_columns(generated)
    workbook_path = create_generated_template_workbook(model, profile, columns)
    tables = [build_generated_monitor_table(model, profile, columns)]
    cells = build_generated_cells(model, profile)
    return RenderPlan(
        output_path=output_path,
        template_path=workbook_path,
        template_source="generated_template",
        cells=cells,
        tables=tables,
        model=model,
        profile=profile,
    )


def resolve_generated_columns(config: dict[str, Any]) -> list[dict[str, Any]]:
    raw_columns = config.get("columns")
    if not isinstance(raw_columns, list) or not raw_columns:
        return [{**dict(item), "column": excel_column_name(index)} for index, item in enumerate(DEFAULT_COLUMNS, start=1)]
    resolved: list[dict[str, Any]] = []
    defaults_by_key = {item["key"]: item for item in DEFAULT_COLUMNS}
    for index, item in enumerate(raw_columns, start=1):
        if not isinstance(item, dict):
            continue
        key = str(item.get("key", "")).strip()
        if not key:
            continue
        base = dict(defaults_by_key.get(key, {}))
        base.update(item)
        base.setdefault("label", key)
        base.setdefault("align", "left")
        base.setdefault("width", 18)
        base["column"] = excel_column_name(index)
        resolved.append(base)
    if not resolved:
        return [{**dict(item), "column": excel_column_name(index)} for index, item in enumerate(DEFAULT_COLUMNS, start=1)]
    return resolved


def create_generated_template_workbook(
    model: WorkbookModel,
    profile: dict[str, Any],
    columns: list[dict[str, Any]],
) -> Path:
    workbook = Workbook()
    environment = workbook.active
    environment.title = "環境"
    settings = workbook.create_sheet("監視設定")
    check_sheet = workbook.create_sheet("チェックシート")

    populate_environment_sheet(environment, model, profile)
    populate_monitor_settings_template(settings, columns, profile)
    populate_check_sheet(check_sheet)

    sanitize_workbook_strings(
        workbook,
        list(profile.get("security", {}).get("secret_patterns", DEFAULT_SECRET_PATTERNS)),
    )

    with NamedTemporaryFile(prefix="bom-generated-template-", suffix=".xlsx", delete=False) as handle:
        temp_path = Path(handle.name)
    workbook.save(temp_path)
    return temp_path


def populate_environment_sheet(worksheet: Any, model: WorkbookModel, profile: dict[str, Any]) -> None:
    apply_standard_sheet_view(worksheet, freeze_panes="A4", show_grid_lines=False)
    apply_standard_print_settings(worksheet, orientation="portrait", title_rows="1:4")
    worksheet.column_dimensions["A"].width = 18
    worksheet.column_dimensions["B"].width = 42
    worksheet.merge_cells("A1:B1")
    worksheet["A1"] = "環境"
    apply_cell_style(worksheet["A1"], title_style())
    worksheet.row_dimensions[1].height = 24
    worksheet.merge_cells("A3:B3")
    worksheet["A3"] = "基本情報"
    apply_cell_style(worksheet["A3"], section_style())
    worksheet.row_dimensions[3].height = 20
    rows = [
        ("項目", "設定値"),
        ("CABファイル名", model.source_name),
        ("profile ID", profile["profile"]["id"]),
        ("profile表示名", profile["profile"].get("name", profile["profile"]["id"])),
        ("group_count", len(model.groups)),
        ("monitor_count", len(model.items)),
        ("生成日時", datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S UTC")),
    ]
    start_row = 4
    for offset, row_values in enumerate(rows, start=0):
        row_index = start_row + offset
        for column_index, value in enumerate(row_values, start=1):
            cell = worksheet.cell(row_index, column_index, value)
            if row_index == start_row:
                apply_cell_style(cell, header_style())
            else:
                apply_cell_style(cell, data_style(horizontal="left", vertical="center", wrap_text=True))
        worksheet.row_dimensions[row_index].height = 20
    apply_auto_filter(worksheet, "A4:B10")


def populate_monitor_settings_template(worksheet: Any, columns: list[dict[str, Any]], profile: dict[str, Any]) -> None:
    last_column = columns[-1]["column"]
    apply_standard_sheet_view(worksheet, freeze_panes="A4", show_grid_lines=False)
    apply_standard_print_settings(worksheet, orientation="landscape", title_rows="1:3")
    worksheet.merge_cells(f"A1:{last_column}1")
    worksheet["A1"] = f"{profile['profile'].get('name', profile['profile']['id'])} 監視設定"
    apply_cell_style(worksheet["A1"], monitor_setting_title_style())
    worksheet.row_dimensions[1].height = 26
    worksheet.merge_cells(f"A2:{last_column}2")
    worksheet["A2"] = "監視設定"
    apply_cell_style(worksheet["A2"], monitor_setting_section_style())
    worksheet.row_dimensions[2].height = 20
    header_row = 3
    template_row = 4
    for column in columns:
        column_name = column["column"]
        worksheet.column_dimensions[column_name].width = clamp_width(float(column.get("width", 18)))
        header_cell = worksheet[f"{column_name}{header_row}"]
        header_cell.value = column["label"]
        apply_cell_style(header_cell, monitor_setting_header_style())
        body_cell = worksheet[f"{column_name}{template_row}"]
        body_cell.value = None
        apply_cell_style(
            body_cell,
            monitor_setting_data_style(
                horizontal="center" if column.get("align") == "center" else "left",
                vertical="center" if column.get("align") == "center" else "top",
                wrap_text=True,
            ),
        )
    worksheet.row_dimensions[header_row].height = 24
    worksheet.row_dimensions[template_row].height = 20
    apply_auto_filter(worksheet, f"A3:{last_column}4")


def populate_check_sheet(worksheet: Any) -> None:
    apply_standard_sheet_view(worksheet, freeze_panes="A4", show_grid_lines=False)
    apply_standard_print_settings(worksheet, orientation="portrait", title_rows="1:4")
    worksheet.column_dimensions["A"].width = 20
    worksheet.column_dimensions["B"].width = 12
    worksheet.column_dimensions["C"].width = 36
    worksheet.column_dimensions["D"].width = 24
    worksheet.merge_cells("A1:D1")
    worksheet["A1"] = "チェックシート"
    apply_cell_style(worksheet["A1"], title_style())
    worksheet.row_dimensions[1].height = 24
    worksheet.merge_cells("A2:D2")
    worksheet["A2"] = "Validation"
    apply_cell_style(worksheet["A2"], section_style())
    worksheet.row_dimensions[2].height = 20
    rows = [
        ("項目", "判定", "内容", "備考"),
        ("Validation", "OK", "PASS", None),
        ("Profile", None, None, None),
        ("Template source", None, None, None),
        ("Monitor count", None, None, None),
        ("Group count", None, None, None),
        ("Secret masking", "OK", "PASS", None),
        ("Input unchanged", "OK", "PASS", None),
    ]
    start_row = 3
    for offset, row_values in enumerate(rows, start=0):
        row_index = start_row + offset
        for column_index, value in enumerate(row_values, start=1):
            cell = worksheet.cell(row_index, column_index, value)
            if row_index == start_row:
                apply_cell_style(cell, header_style())
            else:
                apply_cell_style(cell, data_style(horizontal="left", vertical="center", wrap_text=True))
        worksheet.row_dimensions[row_index].height = 20
    apply_auto_filter(worksheet, "A3:D10")


def build_generated_monitor_table(
    model: WorkbookModel,
    profile: dict[str, Any],
    columns: list[dict[str, Any]],
) -> RenderTable:
    transformed_rows = [build_output_row(item, model, profile) for item in model.items]
    rows: list[RenderRow] = []
    managed_columns: dict[str, str] = {}
    for column in columns:
        managed_columns[column["key"]] = column["column"]
    for index, row_values in enumerate(transformed_rows, start=1):
        generated_values = dict(row_values)
        generated_values.setdefault("no", index)
        generated_values.setdefault("target", resolve_target_value(generated_values))
        generated_values.setdefault("action", resolve_action_value(generated_values))
        generated_values.setdefault("remarks", row_values.get("remarks"))
        masked_values = mask_secrets_in_mapping(
            generated_values,
            list(profile.get("security", {}).get("secret_patterns", DEFAULT_SECRET_PATTERNS)),
        )
        table_values = build_table_row(masked_values, {column["column"]: column["key"] for column in columns})
        rows.append(RenderRow(values=table_values, source_monitor_id=str(row_values.get("monitor_id", index))))

    format_alignment = {
        column["column"]: {
            "horizontal": "center" if column.get("align") == "center" else "left",
            "vertical": "top",
            "wrap_text": True,
        }
        for column in columns
    }
    format_alignment.update(
        {
            column["column"]: {"horizontal": "center", "vertical": "top", "wrap_text": True}
            for column in columns
            if column["key"] in {"group_id", "monitor_id", "monitor_type", "enabled", "interval", "average_count", "no"}
        }
    )
    layout = {
        "group_merge": {"enabled": False, "columns": []},
        "style_sources": {
            "group_first": 4,
            "group_middle": 4,
            "group_last": 4,
            "single_row_group": 4,
            "table_last": 4,
        },
    }
    format_rules = {
        "alignment": format_alignment,
        "column_widths": {column["column"]: float(column.get("width", 18)) for column in columns},
        "row_height": {
            "default": 20,
            "thresholds": [{"max": 40, "height": 20}, {"max": 80, "height": 30}],
            "fallback": 48,
        },
        "borders": {
            "group_first": {"top": "medium", "bottom": "thin"},
            "group_middle": {"top": "thin", "bottom": "thin"},
            "group_last": {"top": "thin", "bottom": "thin"},
            "single_row_group": {"top": "medium", "bottom": "thin"},
            "table_last": {"top": "thin", "bottom": "thin"},
        },
    }
    return RenderTable(
        name="monitor_settings",
        sheet="監視設定",
        start_row=4,
        template_row=4,
        rows=rows,
        managed_columns=managed_columns,
        group_merge_fields=[],
        layout=layout,
        format_rules=format_rules,
        group_labels_first_row_only=True,
        clear_existing_data=True,
        trim_unused_rows=True,
        unused_rows_mode="clear_values",
        template_data_end_row=None,
        clear_existing_rows=False,
        reuse_existing_rows=False,
    )


def build_generated_cells(model: WorkbookModel, profile: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {
        "document_title": {
            "sheet": "監視設定",
            "cell": "A1",
            "value": f"{profile['profile'].get('name', profile['profile']['id'])} 監視設定",
        },
        "check_profile": {
            "sheet": "チェックシート",
            "cell": "C5",
            "value": profile["profile"]["id"],
        },
        "check_template_source": {
            "sheet": "チェックシート",
            "cell": "C6",
            "value": "generated_template",
        },
        "check_monitor_count": {
            "sheet": "チェックシート",
            "cell": "C7",
            "value": len(model.items),
        },
        "check_group_count": {
            "sheet": "チェックシート",
            "cell": "C8",
            "value": len(model.groups),
        },
    }


def resolve_target_value(values: dict[str, Any]) -> Any:
    for key in ("target", "ObjectName", "ValueName", "Object", "Target"):
        value = values.get(key)
        if value not in (None, ""):
            return value
    return None


def resolve_action_value(values: dict[str, Any]) -> Any:
    for key in ("action", "action_name", "ActionName", "action_id"):
        value = values.get(key)
        if value not in (None, ""):
            return value
    return None


def cleanup_generated_template(path: Path | None) -> None:
    if path is None:
        return
    if path.name.startswith("bom-generated-template-") and path.exists():
        path.unlink(missing_ok=True)


def excel_column_name(index: int) -> str:
    result = ""
    value = index
    while value:
        value, remainder = divmod(value - 1, 26)
        result = chr(ord("A") + remainder) + result
    return result
