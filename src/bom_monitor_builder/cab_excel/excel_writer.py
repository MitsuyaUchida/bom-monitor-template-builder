"""Workbook generation with openpyxl."""

from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]

from .formatter import format_enabled, format_interval, format_threshold, sanitize_excel_text
from .models import MonitorGroup, ParsedCab
from .options_parser import parse_options

HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
HEADER_FONT = Font(bold=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def build_workbook(parsed: ParsedCab, *, tool_version: str) -> Workbook:
    """Build the Excel workbook for a parsed CAB."""
    workbook = Workbook()
    cover = workbook.active
    cover.title = "表紙"
    write_cover_sheet(cover, parsed, tool_version=tool_version)

    write_groups_sheet(workbook.create_sheet("監視グループ一覧"), parsed.groups)
    write_items_sheet(workbook.create_sheet("監視項目一覧"), parsed)
    if parsed.actions:
        write_actions_sheet(workbook.create_sheet("Actions"), parsed)
    write_item_details_sheet(workbook.create_sheet("監視項目詳細"), parsed)
    write_xml_sheet(workbook.create_sheet("XML全項目"), parsed)
    write_analysis_sheet(workbook.create_sheet("解析情報"), parsed)
    return workbook



def write_actions_sheet(sheet: Worksheet, parsed: ParsedCab) -> None:
    """Write action XML values separately from monitor rows."""
    keys = sorted({key for action in parsed.actions for key in action.raw_values})
    headers = ["Group folder", "Action XML", "Monitor XML", *keys]
    write_header(sheet, headers)
    for row_index, action in enumerate(parsed.actions, start=2):
        monitor_xml = Path(action.xml_file).name.split("ACT", 1)[0] + ".xml"
        values = [action.group_folder, action.xml_file, monitor_xml, *[action.raw_values.get(key, "") for key in keys]]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, sanitize_excel_text(value))
    finalize_sheet(sheet, freeze="A2", landscape=True, repeat_rows="1:1")

def write_cover_sheet(sheet: Worksheet, parsed: ParsedCab, *, tool_version: str) -> None:
    """Write the cover sheet."""
    rows = [
        ("文書名", "BOM監視設定仕様書"),
        ("元CABファイル名", parsed.source_name),
        ("製品名", parsed.manifest.product),
        ("エクスポート種別", parsed.manifest.export_type),
        ("BOMバージョン", parsed.manifest.bom_version),
        ("監視グループ数", len(parsed.groups)),
        ("監視項目数", len(parsed.items)),
        ("作成日時", parsed.parsed_completed_at),
        ("生成ツールバージョン", tool_version),
    ]
    for row_index, (label, value) in enumerate(rows, start=1):
        sheet.cell(row=row_index, column=1, value=label)
        value_cell = sheet.cell(row=row_index, column=2, value=sanitize_excel_text(value))
        if isinstance(value, datetime):
            value_cell.number_format = "yyyy/mm/dd hh:mm:ss"
        apply_header_style(sheet.cell(row=row_index, column=1))
    finalize_sheet(sheet, freeze="A1", landscape=False, repeat_rows=None)
    set_column_widths(sheet, {1: 24, 2: 40})


def write_groups_sheet(sheet: Worksheet, groups: list[MonitorGroup]) -> None:
    """Write group summary rows."""
    headers = [
        "No.",
        "グループフォルダー",
        "グループ名",
        "有効",
        "コメント",
        "アイコン番号",
        "月曜日",
        "火曜日",
        "水曜日",
        "木曜日",
        "金曜日",
        "土曜日",
        "日曜日",
        "ParentId",
        "Id",
    ]
    write_header(sheet, headers)
    day_keys = ["TTMon", "TTTue", "TTWed", "TTThu", "TTFri", "TTSat", "TTSun"]
    for index, group in enumerate(groups, start=2):
        row = [
            index - 1,
            group.group_folder,
            group.raw_values.get("Name", ""),
            format_enabled(group.typed_values.get("Enabled")),
            group.raw_values.get("Comments", ""),
            group.typed_values.get("IconIndex"),
            *[group.raw_values.get(key, "") for key in day_keys],
            group.raw_values.get("ParentId", ""),
            group.raw_values.get("Id", ""),
        ]
        for column, value in enumerate(row, start=1):
            sheet.cell(index, column, sanitize_excel_text(value))
    finalize_sheet(sheet, freeze="A2", landscape=True, repeat_rows="1:1")
    set_column_widths(sheet, {2: 18, 3: 28, 5: 40})


def write_items_sheet(sheet: Worksheet, parsed: ParsedCab) -> None:
    """Write monitor item summary rows."""
    headers = [
        "No.",
        "グループ名",
        "監視ファイル",
        "監視名",
        "監視タイプ",
        "有効",
        "監視間隔",
        "開始時刻",
        "注意判定",
        "危険判定",
        "表示単位",
        "スケール",
        "実行オブジェクト",
        "値名",
        "コメント",
    ]
    write_header(sheet, headers)
    group_names = {group.group_folder: group.raw_values.get("Name", "") for group in parsed.groups}
    for index, item in enumerate(parsed.items, start=2):
        interval_text, _ = format_interval(
            item.typed_values.get("Interval"),
            item.raw_values.get("IntervalUnit", ""),
        )
        warning_text, _ = format_threshold(
            item.typed_values.get("CmpValueY"),
            item.raw_values.get("CmpMethodY", ""),
        )
        danger_text, _ = format_threshold(
            item.typed_values.get("CmpValueR"),
            item.raw_values.get("CmpMethodR", ""),
        )
        row = [
            index - 1,
            group_names.get(item.group_folder, item.group_folder),
            Path(item.xml_file).name,
            item.raw_values.get("Name", ""),
            item.raw_values.get("Type", ""),
            format_enabled(item.typed_values.get("Enabled")),
            interval_text,
            item.raw_values.get("StartTime", ""),
            warning_text,
            danger_text,
            item.raw_values.get("DisplayUnit", ""),
            item.raw_values.get("Scale", ""),
            item.raw_values.get("ObjectName", ""),
            item.raw_values.get("ValueName", ""),
            item.raw_values.get("Comments", ""),
        ]
        for column, value in enumerate(row, start=1):
            sheet.cell(index, column, sanitize_excel_text(value))
    finalize_sheet(sheet, freeze="A2", landscape=True, repeat_rows="1:1")
    set_column_widths(sheet, {2: 20, 4: 32, 8: 18, 13: 28, 15: 40})


def write_item_details_sheet(sheet: Worksheet, parsed: ParsedCab) -> None:
    """Write monitor item blocks."""
    row = 1
    group_names = {group.group_folder: group.raw_values.get("Name", "") for group in parsed.groups}
    for item_index, item in enumerate(parsed.items, start=1):
        options = parse_options(item.raw_values.get("Options", ""))
        interval_text, _ = format_interval(
            item.typed_values.get("Interval"),
            item.raw_values.get("IntervalUnit", ""),
        )
        warning_text, _ = format_threshold(
            item.typed_values.get("CmpValueY"),
            item.raw_values.get("CmpMethodY", ""),
        )
        danger_text, _ = format_threshold(
            item.typed_values.get("CmpValueR"),
            item.raw_values.get("CmpMethodR", ""),
        )
        entries = [
            ("監視No.", item_index),
            ("グループ名", group_names.get(item.group_folder, item.group_folder)),
            ("監視名", item.raw_values.get("Name", "")),
            ("監視タイプ", item.raw_values.get("Type", "")),
            ("有効", format_enabled(item.typed_values.get("Enabled"))),
            ("開始時刻", item.raw_values.get("StartTime", "")),
            ("監視間隔", interval_text),
            ("CLSID", item.raw_values.get("CLSID", "")),
            ("ObjectName", item.raw_values.get("ObjectName", "")),
            ("ValueName", item.raw_values.get("ValueName", "")),
            ("Options", item.raw_values.get("Options", "")),
            ("RetryInterval", options.retry_interval),
            ("Timeout", options.timeout),
            ("実行スクリプト", options.script_path),
            ("注意比較方法", item.raw_values.get("CmpMethodY", "")),
            ("注意しきい値", warning_text),
            ("危険比較方法", item.raw_values.get("CmpMethodR", "")),
            ("危険しきい値", danger_text),
            ("表示単位", item.raw_values.get("DisplayUnit", "")),
            ("スケール", item.raw_values.get("Scale", "")),
            ("コメント", item.raw_values.get("Comments", "")),
            ("XMLファイルパス", item.xml_file),
        ]
        for label, value in entries:
            sheet.cell(row=row, column=1, value=label)
            sheet.cell(row=row, column=2, value=sanitize_excel_text(value))
            apply_header_style(sheet.cell(row=row, column=1))
            row += 1
        row += 1
    finalize_sheet(sheet, freeze="A1", landscape=False, repeat_rows=None)
    set_column_widths(sheet, {1: 20, 2: 60})


def write_xml_sheet(sheet: Worksheet, parsed: ParsedCab) -> None:
    """Write flattened XML values."""
    headers = ["種別", "グループフォルダー", "XMLファイル", "監視名またはグループ名", "要素名", "値"]
    write_header(sheet, headers)
    row = 2
    for record in [*parsed.groups, *parsed.items]:
        for key, value in record.raw_values.items():
            sheet.cell(row, 1, record.kind)
            sheet.cell(row, 2, record.group_folder)
            sheet.cell(row, 3, record.xml_file)
            sheet.cell(row, 4, sanitize_excel_text(record.display_name))
            sheet.cell(row, 5, key)
            sheet.cell(row, 6, sanitize_excel_text(value))
            row += 1
    finalize_sheet(sheet, freeze="A2", landscape=True, repeat_rows="1:1")
    set_column_widths(sheet, {3: 30, 4: 28, 6: 50})


def write_analysis_sheet(sheet: Worksheet, parsed: ParsedCab) -> None:
    """Write manifest values, processing details, and warnings."""
    row = 1
    sheet.cell(row, 1, "区分")
    sheet.cell(row, 2, "キー")
    sheet.cell(row, 3, "値")
    write_header_style_row(sheet, 1, 3)
    row += 1

    for key, manifest_value in parsed.manifest.values.items():
        sheet.cell(row, 1, "MANIFEST")
        sheet.cell(row, 2, key)
        sheet.cell(row, 3, sanitize_excel_text(manifest_value))
        row += 1

    analysis_rows: list[tuple[str, str, object]] = [
        ("解析", "CAB展開方式", parsed.extraction_method),
        ("解析", "入力SHA-256", parsed.input_sha256),
        ("解析", "解析開始日時", parsed.parsed_started_at),
        ("解析", "解析完了日時", parsed.parsed_completed_at),
        ("解析", "認識監視タイプ", ", ".join(parsed.recognized_monitor_types)),
        ("解析", "未知比較演算子", ", ".join(parsed.unknown_comparison_values)),
        ("解析", "未知間隔単位", ", ".join(parsed.unknown_interval_units)),
    ]
    for section, analysis_key, value in analysis_rows:
        sheet.cell(row, 1, section)
        sheet.cell(row, 2, analysis_key)
        value_cell = sheet.cell(row, 3, sanitize_excel_text(value))
        if isinstance(value, datetime):
            value_cell.number_format = "yyyy/mm/dd hh:mm:ss"
        row += 1

    for warning in parsed.warnings:
        sheet.cell(row, 1, "警告")
        sheet.cell(row, 2, warning.code)
        message = warning.message if warning.path is None else f"{warning.message} ({warning.path})"
        sheet.cell(row, 3, sanitize_excel_text(message))
        row += 1

    finalize_sheet(sheet, freeze="A2", landscape=False, repeat_rows="1:1")
    set_column_widths(sheet, {1: 12, 2: 22, 3: 60})


def save_workbook(parsed: ParsedCab, output_path: Path, *, tool_version: str) -> None:
    """Save workbook to disk."""
    workbook = build_workbook(parsed, tool_version=tool_version)
    workbook.save(output_path)


def write_header(sheet: Worksheet, headers: list[str]) -> None:
    """Write a header row."""
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(1, column, header)
        apply_header_style(cell)


def write_header_style_row(sheet: Worksheet, row: int, column_count: int) -> None:
    """Apply header style to a row."""
    for column in range(1, column_count + 1):
        apply_header_style(sheet.cell(row, column))


def apply_header_style(cell: Any) -> None:
    """Apply a consistent header style."""
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = THIN_BORDER


def finalize_sheet(
    sheet: Worksheet,
    *,
    freeze: str,
    landscape: bool,
    repeat_rows: str | None,
) -> None:
    """Apply generic sheet formatting."""
    sheet.freeze_panes = freeze
    if sheet.max_row >= 1 and sheet.max_column >= 1:
        sheet.auto_filter.ref = sheet.dimensions
    if repeat_rows is not None:
        sheet.print_title_rows = repeat_rows
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def set_column_widths(sheet: Worksheet, presets: dict[int, int]) -> None:
    """Best-effort column width adjustment with a maximum width."""
    widths: dict[int, int] = defaultdict(int)
    for row in sheet.iter_rows():
        for cell in row:
            value = "" if cell.value is None else str(cell.value)
            widths[cell.column] = max(widths[cell.column], min(max(len(value) + 2, 10), 60))
    for column, width in presets.items():
        widths[column] = max(widths[column], min(width, 60))
    for column_index, width in widths.items():
        sheet.column_dimensions[get_column_letter(column_index)].width = min(width, 60)
