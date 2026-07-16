from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook  # type: ignore[import-untyped]

from bom_monitor_builder.knowledge.output import ensure_within_root, write_csv, write_json
from bom_monitor_builder.template.models import ScanResult


def write_scan_outputs(scan_result: ScanResult, output_root: Path) -> None:
    inventory_root = ensure_within_root(output_root, "inventory")
    reports_root = ensure_within_root(output_root, "reports")
    inventory_root.mkdir(parents=True, exist_ok=True)
    reports_root.mkdir(parents=True, exist_ok=True)

    write_json(ensure_within_root(output_root, "inventory/directory_tree.json"), scan_result.directory_tree)
    write_json(
        ensure_within_root(output_root, "inventory/files.json"),
        {"files": [entry.to_dict() for entry in scan_result.files]},
    )
    write_json(
        ensure_within_root(output_root, "inventory/categories.json"),
        {"categories": [entry.to_dict() for entry in scan_result.categories]},
    )
    write_json(
        ensure_within_root(output_root, "inventory/templates.json"),
        {"templates": [entry.to_dict() for entry in scan_result.templates]},
    )
    write_json(
        ensure_within_root(output_root, "inventory/extensions.json"),
        {"extensions": [entry.to_dict() for entry in scan_result.extensions]},
    )
    write_json(
        ensure_within_root(output_root, "inventory/summary.json"),
        scan_result.summary.to_dict(),
    )

    write_csv(
        ensure_within_root(output_root, "reports/category_inventory.csv"),
        fieldnames=[
            "category_id",
            "category_name",
            "directory",
            "file_count",
            "cab_count",
            "html_count",
        ],
        rows=[entry.to_dict() for entry in scan_result.categories],
    )
    write_csv(
        ensure_within_root(output_root, "reports/template_inventory.csv"),
        fieldnames=[
            "category_id",
            "category_name",
            "directory",
            "template_name_candidate",
            "cab",
            "html",
            "same_name_pair",
            "html_only",
            "cab_only",
        ],
        rows=[entry.to_dict() for entry in scan_result.templates],
    )
    write_csv(
        ensure_within_root(output_root, "reports/file_inventory.csv"),
        fieldnames=[
            "relative_path",
            "category_id",
            "category_name",
            "directory",
            "file_name",
            "template_name_candidate",
            "extension",
            "file_type",
            "size",
            "modified_at",
            "sha256",
            "depth",
            "is_japanese_named",
        ],
        rows=[entry.to_dict() for entry in scan_result.files],
    )

    write_excel_summary(scan_result, ensure_within_root(output_root, "reports/template_summary.xlsx"))


def write_excel_summary(scan_result: ScanResult, output_path: Path) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    for row in build_summary_rows(scan_result.summary.to_dict()):
        summary_sheet.append(row)

    categories_sheet = workbook.create_sheet("Categories")
    append_table(categories_sheet, [entry.to_dict() for entry in scan_result.categories])

    templates_sheet = workbook.create_sheet("Templates")
    append_table(templates_sheet, [entry.to_dict() for entry in scan_result.templates])

    files_sheet = workbook.create_sheet("Files")
    append_table(files_sheet, [entry.to_dict() for entry in scan_result.files])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def build_summary_rows(payload: dict[str, Any], prefix: str = "") -> list[list[Any]]:
    rows: list[list[Any]] = []
    for key, value in payload.items():
        label = f"{prefix}{key}" if not prefix else f"{prefix}.{key}"
        if isinstance(value, dict):
            rows.extend(build_summary_rows(value, label))
        elif isinstance(value, list):
            rows.append([label, ", ".join(str(item) for item in value)])
        else:
            rows.append([label, value])
    return rows


def append_table(worksheet: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        worksheet.append(["no_data"])
        return
    headers = list(rows[0].keys())
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header) for header in headers])
